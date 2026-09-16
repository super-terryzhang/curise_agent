"""Generate the master-data product upload template (.xlsx).

Run this script when `_HEADER_ALIASES` in
`domains.masterdata.upload.service` changes — it regenerates the canonical
template at `v3_backend/static/templates/product_upload_template.xlsx`.

Self-consistency tests (`test_v2/section_6_masterdata/test_upload_template.py`)
reload the generated file through `parse_excel` to catch the "I forgot
to regenerate" drift class in CI.

Design philosophy (revised 2026-05-14 after the v18 review):

The template mirrors the SINGLE-record edit form in `/dashboard/data`:
that's the user's mental model of "what a product is". Anything less and
users hit the "I have 1000 SKUs but can only batch-edit 5 fields" wall.
So this template ships all 21 supported Product fields, and the
backend pipeline accepts every one of them.

FK fields (category / supplier / country / port) are filled by **NAME**,
not by id. Two reasons:
  1. Names are what users read in the existing UI dropdowns.
  2. Users typing "Australia" is faster + safer than typing `42` and
     getting it wrong.
The resolver looks each name up in masterdata at parse time; misses
become row-level validation errors with a clear message ("country 'X'
not found").

Cells are formatted to defend against the most common Excel footguns:
  - `product_code` cell type = "@" (text) — keeps "00100" as 5 chars.
  - `price` = "#,##0.00" — visual cue this is numeric.
  - all effective-date columns = "yyyy-mm-dd" — narrows the input.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

# ─── Field catalog ───────────────────────────────────────────────
#
# (canonical_name, required, display_zh, example, format_hint, notes)
# Order = column order in the data sheet. Required field first; then
# identity (code); then taxonomy (FKs); then commercial (price/unit);
# then metadata (origin, dates).

REQUIRED_COLUMNS = [
    ("product_name", True, "产品名称（英文）", "Apple - Red Delicious 125ct", "文本",
     "对应 UI 的「英文品名」。必填。"),
    # 2026-05-27 strict identity contract: same product code at different
    # ports/countries = different records. Without these fields the matcher
    # cannot deterministically identify a row → row-level error.
    ("country", True, "国家（名称）", "USA",
     "国家名称（不是 id）",
     "对应 UI 的「国家」下拉。**必填**——产品的身份键之一。按 countries.name 查找。注意这是**采购国**，不是原产地。"),
    ("port", True, "港口（名称）", "Yokohama",
     "港口名称（不是 id）",
     "对应 UI 的「港口」下拉。**必填**——产品的身份键之一。按 ports.name 查找。同一产品在不同港口 = 不同记录。"),
]

OPTIONAL_COLUMNS = [
    # Identity
    ("product_code", False, "产品代码 / SKU", "FRT-APL-RED",
     "文本（保留前导零）",
     "对应 UI 的「商品代码」。系统按 (code, country, port) 三元组匹配现有产品。空时按 (country, name, port) 匹配。"),
    ("product_name_jp", False, "产品名称（日文）", "赤りんご デリシャス",
     "文本",
     "对应 UI 的「日文品名」。空表示不更新已有值。"),
    ("brand", False, "品牌", "Sunkist",
     "文本",
     "对应 UI 的「品牌」。"),
    # FK by NAME
    ("category", False, "类别（名称）", "FRUIT",
     "类别名称（不是 id）",
     "对应 UI 的「类别」下拉。系统按 categories 表的 `name` 字段查找；找不到 → 行级错误。可在「数据管理」→「类别」Tab 查看 / 新增。"),
    ("supplier", False, "供应商（名称）", "Sunkist Growers Inc.",
     "供应商名称（不是 id）",
     "对应 UI 的「供应商」下拉。按 suppliers.name 查找；找不到 → 行级错误。"),
    # Commercial
    ("price", False, "采购价", 850.00,
     "数字（不带 ¥ $ 符号）",
     "对应 UI 的「采购价」。我方付给供应商的价格。0 是有效值（系统会保留）。"),
    ("purchase_price_effective_from", False, "采购价有效开始日期", "2026-01-01",
     "日期 YYYY-MM-DD",
     "采购价开始生效的日期。新产品空 = 不限制起始；更新时空 = 保留旧值。"),
    ("purchase_price_effective_to", False, "采购价有效结束日期", "2026-06-30",
     "日期 YYYY-MM-DD",
     "采购价有效的最后一天。新产品空 = 不限制结束；更新时空 = 保留旧值。"),
    ("contract_price", False, "卖价", 1050.00,
     "数字（不带 ¥ $ 符号）",
     "对应 UI 的「卖价」。我方卖给邮轮客户的价格 —— 财务对比 PO 单价的基线。"
     "上传时也接受别名「合同价」/「合同卖价」/「合約価格」/「契約価格」（向后兼容）。"),
    ("selling_price_effective_from", False, "卖价有效开始日期", "2026-01-01",
     "日期 YYYY-MM-DD",
     "卖价开始生效的日期。新产品空 = 不限制起始；更新时空 = 保留旧值。"),
    ("selling_price_effective_to", False, "卖价有效结束日期", "2026-12-31",
     "日期 YYYY-MM-DD",
     "卖价有效的最后一天。新产品空 = 不限制结束；更新时空 = 保留旧值。"),
    ("currency", False, "币种", "USD",
     "ISO 4217 三字母代码",
     "对应 UI 的「币种」。建议：USD / JPY / EUR / AUD / CNY 等。"),
    ("unit", False, "计量单位", "CT",
     "短代码",
     "对应 UI 的「单位」。常用：KG / EA / L / CT / G / PCS / BOX。"),
    ("unit_size", False, "单位规格", "10kg",
     "自由文本",
     "对应 UI 的「单位规格」。"),
    ("pack_size", False, "包装规格", "40LB/CT",
     "自由文本",
     "对应 UI 的「包装规格」。"),
    # Origin
    ("country_of_origin", False, "原产地", "Washington, USA",
     "自由文本",
     "对应 UI 的「原产地」。**字符串字段**（不是 FK）—— 跟「country」（采购国）是两个独立字段。"),
    # Effective period
    ("effective_from", False, "产品有效开始日期", "2026-01-01",
     "日期 YYYY-MM-DD",
     "控制产品本身是否可用于匹配，不代表采购价或卖价有效期。新产品空 = 不限制起始；更新时空 = 保留旧值。"),
    ("effective_to", False, "产品有效结束日期", "2026-12-31",
     "日期 YYYY-MM-DD",
     "控制产品本身是否可用于匹配，不代表采购价或卖价有效期。新产品空 = 不限制结束；更新时空 = 保留旧值。"),
]

ALL_COLUMNS = REQUIRED_COLUMNS + OPTIONAL_COLUMNS


def _row_dict(values: list) -> dict[str, object]:
    """Zip a value list with the canonical names so example rows stay in
    sync if column order changes."""
    return {col[0]: v for col, v in zip(ALL_COLUMNS, values, strict=False)}


# Three example rows covering distinct gotchas:
#   row 0: standard happy path with most fields filled
#   row 1: leading-zero SKU (text format must survive Excel auto-coerce)
#   row 2: minimal — only the THREE required fields (name + country + port)
#          shows the floor under the strict identity contract.
# Column order matches REQUIRED_COLUMNS + OPTIONAL_COLUMNS:
#   product_name, country, port, product_code, product_name_jp, brand,
#   category, supplier, price, purchase-price dates, contract_price,
#   selling-price dates, currency, unit, unit_size, pack_size,
#   country_of_origin, product effective_from, product effective_to
EXAMPLE_ROWS: list[dict[str, object]] = [
    _row_dict([
        "Apple - Red Delicious 125ct",         # product_name *
        "USA",                                 # country *
        "Yokohama",                            # port *
        "FRT-APL-RED",                         # product_code
        "赤りんご デリシャス",                   # product_name_jp
        "Sunkist",                             # brand
        "FRUIT",                               # category
        "Sunkist Growers Inc.",                # supplier
        850.00,                                # price (采购价)
        datetime(2026, 1, 1),                  # purchase price from
        datetime(2026, 6, 30),                 # purchase price to
        1050.00,                               # contract_price (卖价)
        datetime(2026, 1, 1),                  # selling price from
        datetime(2026, 12, 31),                # selling price to
        "USD",                                 # currency
        "CT",                                  # unit
        "40LB/CT",                             # unit_size
        "125CT/CTN",                           # pack_size
        "Washington, USA",                     # country_of_origin
        datetime(2026, 1, 1),                  # effective_from
        datetime(2026, 12, 31),                # effective_to
    ]),
    _row_dict([
        "Yogurt - Strawberry 75G",             # product_name *
        "Japan",                               # country *
        "Yokohama",                            # port *
        "00100",                               # product_code (leading zeros!)
        "イチゴヨーグルト 75G",                  # product_name_jp
        "Morinaga",                            # brand
        "DAIRY",                               # category
        "Morinaga Milk Industry Co., Ltd.",    # supplier
        76.46,                                 # price (采购价)
        None,                                  # purchase price from
        None,                                  # purchase price to
        95.00,                                 # contract_price (卖价)
        None,                                  # selling price from
        None,                                  # selling price to
        "JPY",                                 # currency
        "CT",                                  # unit
        "75G×4×6PACK",                         # unit_size
        "24CT/CTN",                            # pack_size
        "Japan",                               # country_of_origin
        None,                                  # effective_from (blank ok)
        None,                                  # effective_to
    ]),
    _row_dict([
        "Beef Tenderloin Grade A",             # product_name * (minimum)
        "USA",                                 # country * (required even in minimal row)
        "Yokohama",                            # port * (required even in minimal row)
        None, None, None, None, None, None, None, None, None,
        None, None, None, None, None, None, None, None, None,
    ]),
]


# ─── Visual styling ───────────────────────────────────────────────

_REQUIRED_FILL = PatternFill("solid", fgColor="DBEAFE")   # light blue
_OPTIONAL_FILL = PatternFill("solid", fgColor="F3F4F6")   # light grey
_HEADER_FONT = Font(bold=True, size=11)
_THIN_BORDER = Border(
    left=Side(style="thin", color="D1D5DB"),
    right=Side(style="thin", color="D1D5DB"),
    top=Side(style="thin", color="D1D5DB"),
    bottom=Side(style="thin", color="D1D5DB"),
)


def _build_data_sheet(ws: Worksheet) -> None:
    ws.title = "产品数据"

    # Header row — canonical names. Required columns get a blue fill and
    # an asterisk in the display label (rendered via cell comment).
    for idx, (canonical, required, display, _example, _fmt, _notes) in enumerate(
        ALL_COLUMNS, start=1
    ):
        cell = ws.cell(row=1, column=idx, value=canonical)
        cell.font = _HEADER_FONT
        cell.fill = _REQUIRED_FILL if required else _OPTIONAL_FILL
        cell.border = _THIN_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
        # Hover tooltip carries the human-readable label.
        marker = " *" if required else ""
        cell.comment = Comment(
            f"{display}{marker}\n字段名: {canonical}\n示例: {_example}",
            "system",
        )

    # Example rows.
    for r_idx, row_data in enumerate(EXAMPLE_ROWS, start=2):
        for c, (canonical, _req, _disp, _ex, _fmt_hint, _notes) in enumerate(
            ALL_COLUMNS, start=1
        ):
            value = row_data.get(canonical)
            cell = ws.cell(row=r_idx, column=c, value=value)
            cell.border = _THIN_BORDER
            # Per-field format hints.
            if canonical == "product_code":
                cell.number_format = "@"  # text — preserve leading zeros
            elif canonical in ("price", "contract_price"):
                cell.number_format = "#,##0.00"
            elif canonical in (
                "effective_from", "effective_to",
                "purchase_price_effective_from", "purchase_price_effective_to",
                "selling_price_effective_from", "selling_price_effective_to",
            ):
                cell.number_format = "yyyy-mm-dd"

    # Column widths chosen by hand for readability — long display fields
    # (product_name, supplier) get more room.
    widths = {
        "product_name": 32,
        "product_code": 16,
        "product_name_jp": 24,
        "brand": 14,
        "category": 14,
        "supplier": 28,
        "country": 12,
        "port": 12,
        "price": 10,
        "purchase_price_effective_from": 31,
        "purchase_price_effective_to": 29,
        "contract_price": 18,
        "selling_price_effective_from": 30,
        "selling_price_effective_to": 28,
        "currency": 10,
        "unit": 8,
        "unit_size": 14,
        "pack_size": 16,
        "country_of_origin": 20,
        "effective_from": 14,
        "effective_to": 14,
    }
    for idx, (canonical, *_rest) in enumerate(ALL_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(canonical, 16)

    # Freeze header row.
    ws.freeze_panes = "A2"


def _build_instructions_sheet(ws: Worksheet) -> None:
    ws.title = "使用说明"

    # Title banner.
    ws["A1"] = "产品上传模板 · 使用说明"
    ws["A1"].font = Font(bold=True, size=16)
    ws.merge_cells("A1:E1")

    intro = [
        "",
        "本模板用于通过 AI 助手批量导入 / 更新产品数据（.xlsx）。请先删除 3 行示例，再填写真实产品。",
        "填写「产品数据」工作表后保存，把文件拖入 AI 助手聊天框即可上传。",
        "",
        "关键约定：",
        "  · 标记必填的列必须填，其他列可以留空（留空表示不更新该字段）。",
        "  · 外键字段（类别 / 供应商 / 国家 / 港口）请填**名称**，不是 ID。",
        "  · 系统会按现有 masterdata 的名字精确匹配（不区分大小写）。",
        "  · 名字找不到时，那一行会被标记为错误，不会破坏其他行。",
        "  · 必填：product_name、country、port。先按代码 + 国家 + 港口，再按英文品名 + 国家 + 港口精确匹配。",
        "  · 已匹配产品的 product_name / product_code 不会被改（防误改身份）。",
        "  · 匹配不到会新增产品；空白保留旧值，不能用空白清空价格。提交可能部分成功，请核对错误行。",
        "",
        "—— 各字段详细说明 ——",
    ]
    for i, line in enumerate(intro, start=2):
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=5)
        ws.cell(row=i, column=1, value=line).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 30

    header_row = 2 + len(intro)
    headers = ["列名 (Excel 表头)", "必填?", "示例", "格式 / 可选值", "说明"]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=c, value=h)
        cell.font = _HEADER_FONT
        cell.fill = _REQUIRED_FILL
        cell.border = _THIN_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Field reference table — one row per column in ALL_COLUMNS.
    for i, (canonical, required, _display, example, fmt_hint, notes) in enumerate(
        ALL_COLUMNS, start=header_row + 1
    ):
        required_str = "✅ 必填" if required else "可选"
        ws.cell(row=i, column=1, value=canonical).font = Font(name="Menlo", size=10)
        ws.cell(row=i, column=2, value=required_str)
        ws.cell(row=i, column=3, value=str(example))
        ws.cell(row=i, column=4, value=fmt_hint)
        ws.cell(row=i, column=5, value=notes)
        for c in range(1, 6):
            cell = ws.cell(row=i, column=c)
            cell.border = _THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[i].height = 60

    # Notes section.
    notes_start = header_row + 1 + len(ALL_COLUMNS) + 2
    notes_lines = [
        "—— 外键字段填名字（不是 ID）—— ",
        "类别 / 供应商 / 国家 / 港口 这 4 个字段，请填**名称**：",
        "  · 类别 → 在「数据管理」→「类别」Tab 看完整列表（例：FRUIT, DAIRY, MEAT）",
        "  · 供应商 → 「供应商」Tab（例：Sunkist Growers Inc.）",
        "  · 国家 → 「国家」Tab（例：USA, Japan, Australia）",
        "  · 港口 → 「港口」Tab（例：Yokohama, Hong Kong）",
        "系统会按这些表的 `name` 字段精确匹配，不区分大小写，自动忽略首尾空格。",
        "找不到名字 → 该行会被标记错误，AI 助手会告诉你哪些名字写错了。",
        "",
        "—— 常见错误 ——",
        "1. 把 product_code 改成数字格式 → 00100 变成 100 → 系统找不到已有产品。"
        "用文本格式（模板已经设好）。",
        "2. 采购价 price、卖价 contract_price 填非负数字；超过两位小数按四舍五入处理，0 有效；公式请先粘贴为数值。",
        "3. 国家填 \"美国\" 但 countries 表里只有 \"USA\" → 不命中。先去「数据管理」",
        "查看真实存在的名字。",
        "4. 第 1 行表头改名 / 删列 → 系统识别不到字段。建议保留规范表头；必填列不能删除。",
        "",
        "—— 日期格式 ——",
        "采购价、卖价和产品本身的所有有效日期列接受：",
        "  · 2026-01-01 （推荐）",
        "  · 2026/01/01",
        "  · 2026.01.01",
        "Excel 把日期单元格存成日期类型也可以，系统会自动转换。",
        "采购价和卖价各自的开始日期不能晚于各自的结束日期。",
        "effective_from / effective_to 是产品整体有效期，不替代两种价格的有效期。",
        "",
        "—— 系统不批量管理的字段 ——",
        "下列字段不会被本模板更新，请通过单条编辑修改：",
        "  · status（启用 / 停用）",
        "  · 历史价格记录（系统会保留 ChangeLog，不在本模板里管）",
        "",
        "—— 如何修正错误 ——",
        "AI 助手在 preview 阶段会告诉你每一行的错误。可以：",
        "  · 修正 Excel 重新上传",
        "  · 或者将可选字段留空后重新上传（必填字段不能留空；空白保留旧值）",
        "  · 或者输入「取消」放弃这个 batch",
    ]
    for i, line in enumerate(notes_lines, start=notes_start):
        cell = ws.cell(row=i, column=1, value=line)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 30
        if line.startswith("—— "):
            cell.font = Font(bold=True)
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=5)

    # Column widths.
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 70


def build_template() -> Workbook:
    wb = Workbook()
    _build_data_sheet(wb.active)
    _build_instructions_sheet(wb.create_sheet())
    return wb


def main() -> Path:
    out_dir = Path(__file__).resolve().parent.parent / "static" / "templates"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / "product_upload_template.xlsx"
    wb = build_template()
    wb.save(out_path)
    print(f"wrote {out_path} ({len(ALL_COLUMNS)} columns)")
    return out_path


if __name__ == "__main__":
    main()
