"""P&L Excel exporter — turns a `Pnl` dict (from `compute_pnl().to_dict()`)
into an .xlsx file bytes blob.

Three sheets:
  1. **Summary** — the equation strip rendered as labelled rows
  2. **Cost Items** — user-entered expenses with native + display amounts
  3. **Products** — per-line P&L (the bottom table in the UI)

Why a standalone module: same Pnl shape can be exported by the HTTP
endpoint (download button) or by an agent tool ("export this order's
P&L to Excel and give me the file URL") without one consumer becoming
the other's dependency.
"""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


_HEADER_FILL = PatternFill("solid", fgColor="305496")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_TOTAL_FILL = PatternFill("solid", fgColor="D9E1F2")
_TOTAL_FONT = Font(bold=True)
_POSITIVE_FONT = Font(color="0F7A2F", bold=True)
_NEGATIVE_FONT = Font(color="C00000", bold=True)


def export_pnl_xlsx(pnl_dict: dict[str, Any]) -> bytes:
    """Build a P&L workbook from a `compute_pnl().to_dict()` payload."""
    wb = Workbook()
    _build_summary_sheet(wb.active, pnl_dict)
    _build_cost_items_sheet(wb.create_sheet("成本费用"), pnl_dict)
    _build_products_sheet(wb.create_sheet("产品明细"), pnl_dict)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _build_summary_sheet(ws: Any, pnl: dict[str, Any]) -> None:
    ws.title = "利润摘要"
    meta = pnl.get("meta") or {}
    summary = pnl["summary"]
    display = pnl["display_currency"]

    rows: list[tuple[str, Any]] = [
        ("订单号", meta.get("po_number") or ""),
        ("船名", meta.get("ship_name") or ""),
        ("交货日期", meta.get("delivery_date") or ""),
        ("状态", meta.get("status") or ""),
        ("显示币种", display),
        ("订单原币种", pnl.get("order_currency")),
        ("税率", f"{pnl['tax_rate'] * 100:.2f}%"),
        ("", ""),
        ("产品总金额（收入）", summary["product_revenue"]),
        ("产品成本", summary["product_cost"]),
        ("其他费用合计", summary["extra_costs_total"]),
        ("总成本（产品 + 其他）", summary["total_cost"]),
        ("毛利润", summary["gross_profit"]),
        ("毛利率 (%)", summary["gross_margin"]),
        ("税费", summary["tax_amount"]),
        ("净利润", summary["net_profit"]),
        ("净利率 (%)", summary["net_margin"]),
    ]
    for i, (label, value) in enumerate(rows, start=1):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=value)
        # Highlight the net_profit row in green/red depending on sign.
        if label == "净利润":
            ws.cell(row=i, column=2).font = (
                _POSITIVE_FONT if isinstance(value, (int, float)) and value >= 0 else _NEGATIVE_FONT
            )
        if label.endswith("（收入）") or label.startswith("总成本"):
            ws.cell(row=i, column=1).font = _TOTAL_FONT
            ws.cell(row=i, column=2).font = _TOTAL_FONT

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 22

    # Footer note for warnings, if any.
    warnings = pnl.get("warnings") or []
    if warnings:
        warn_row = len(rows) + 2
        ws.cell(row=warn_row, column=1, value="警告").font = Font(bold=True, color="C00000")
        for j, w in enumerate(warnings, start=1):
            ws.cell(row=warn_row + j, column=1, value=w)


def _build_cost_items_sheet(ws: Any, pnl: dict[str, Any]) -> None:
    display = pnl["display_currency"]
    headers = ["#", "类别", "原币种", "原金额", f"折算 ({display})", "备注", "FX 状态"]
    for col_idx, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col_idx, value=h)
        c.fill = _HEADER_FILL
        c.font = _HEADER_FONT
        c.alignment = Alignment(horizontal="center")

    items = pnl.get("cost_items") or []
    for i, item in enumerate(items, start=1):
        ws.cell(row=i + 1, column=1, value=i)
        ws.cell(row=i + 1, column=2, value=item["category"])
        ws.cell(row=i + 1, column=3, value=item["currency_original"])
        ws.cell(row=i + 1, column=4, value=item["amount_original"])
        ws.cell(row=i + 1, column=5, value=item["amount_display"])
        ws.cell(row=i + 1, column=6, value=item.get("notes") or "")
        ws.cell(row=i + 1, column=7, value="OK" if item["fx_ok"] else "未折算")

    # Total row
    total_row = len(items) + 2
    ws.cell(row=total_row, column=1, value="合计").font = _TOTAL_FONT
    ws.cell(row=total_row, column=5, value=pnl["summary"]["extra_costs_total"]).font = _TOTAL_FONT
    for col in range(1, len(headers) + 1):
        ws.cell(row=total_row, column=col).fill = _TOTAL_FILL

    widths = [5, 18, 8, 14, 16, 30, 10]
    for col_idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = w


def _build_products_sheet(ws: Any, pnl: dict[str, Any]) -> None:
    display = pnl["display_currency"]
    headers = [
        "#",
        "产品代码",
        "产品名称",
        "数量",
        f"单价 ({display})",
        f"总收入 ({display})",
        f"单位成本 ({display})",
        f"总成本 ({display})",
        f"利润 ({display})",
        "利润率 (%)",
        "供应商 ID",
        "匹配状态",
    ]
    for col_idx, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col_idx, value=h)
        c.fill = _HEADER_FILL
        c.font = _HEADER_FONT
        c.alignment = Alignment(horizontal="center")

    lines = pnl.get("product_lines") or []
    for i, p in enumerate(lines, start=1):
        ws.cell(row=i + 1, column=1, value=i)
        ws.cell(row=i + 1, column=2, value=p.get("product_code") or "")
        ws.cell(row=i + 1, column=3, value=p.get("product_name") or "")
        ws.cell(row=i + 1, column=4, value=p["quantity"])
        ws.cell(row=i + 1, column=5, value=p["unit_price"])
        ws.cell(row=i + 1, column=6, value=p["revenue"])
        ws.cell(row=i + 1, column=7, value=p.get("unit_cost") or "")
        ws.cell(row=i + 1, column=8, value=p["cost"])
        ws.cell(row=i + 1, column=9, value=p["profit"])
        ws.cell(row=i + 1, column=10, value=p["margin"])
        ws.cell(row=i + 1, column=11, value=p.get("supplier_id") or "")
        ws.cell(row=i + 1, column=12, value="已匹配" if p["matched"] else "未匹配")

    total_row = len(lines) + 2
    summary = pnl["summary"]
    ws.cell(row=total_row, column=1, value="合计").font = _TOTAL_FONT
    ws.cell(row=total_row, column=6, value=summary["product_revenue"]).font = _TOTAL_FONT
    ws.cell(row=total_row, column=8, value=summary["product_cost"]).font = _TOTAL_FONT
    ws.cell(row=total_row, column=9, value=summary["gross_profit"]).font = _TOTAL_FONT
    ws.cell(row=total_row, column=10, value=summary["gross_margin"]).font = _TOTAL_FONT
    for col in range(1, len(headers) + 1):
        ws.cell(row=total_row, column=col).fill = _TOTAL_FILL

    widths = [4, 16, 38, 8, 12, 14, 14, 14, 14, 10, 12, 10]
    for col_idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = w
