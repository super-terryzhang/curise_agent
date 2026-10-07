"""Bounded workbook parsing into immutable source-coordinate staging rows."""

from __future__ import annotations

import hashlib
import json
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from pydantic import ValidationError as PydanticValidationError
from sqlalchemy.orm import Session

from .catalog import load_product_catalog
from .contracts import WorkbookManifest
from .models import ImportBatch, ImportRow
from .template import (
    DATA_START_ROW,
    HEADER_ROW,
    PRICE_SHEET,
    PRODUCT_SHEET,
    SYSTEM_SHEET,
    TECH_PRICE_HEADERS,
    TECH_PRODUCT_HEADERS,
)

MAX_WORKBOOK_BYTES = 8 * 1024 * 1024
SHEETS = {PRODUCT_SHEET: "products", PRICE_SHEET: "prices"}
TECH_KEYS = {
    "__产品ID": "__product_id",
    "__产品版本": "__product_revision",
    "__扩展版本": "__extension_revision",
    "__价格区间ID": "__period_id",
    "__价格区间版本": "__period_revision",
}


class UnsupportedContract(ValueError):
    pass


def _issue(code: str, message: str, *, sheet=None, row=None, field=None) -> dict:
    issue = {"severity": "block", "code": code, "message": message}
    if sheet is not None:
        issue["sheet"] = sheet
    if row is not None:
        issue["row"] = row
    if field is not None:
        issue["field"] = field
    return issue


def _json_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    return value


def _new_batch(db, file_bytes, filename, user_id, *, manifest=None) -> ImportBatch:
    catalog = load_product_catalog(db)
    batch = ImportBatch(
        user_id=user_id,
        filename=filename[:500],
        file_sha256=hashlib.sha256(file_bytes).hexdigest(),
        target_table="products",
        contract_version=(manifest.contract_version if manifest else 1),
        product_schema_version=(
            manifest.product_schema_version if manifest else catalog.table.schema_version
        ),
        field_manifest=(manifest.model_dump(mode="json") if manifest else {}),
        status="uploaded",
    )
    db.add(batch)
    db.flush()
    return batch


def _finish_failed(db, batch, issue, message):
    batch.status = "failed"
    batch.error_message = message
    batch.block_count = 1
    batch.result = {"issues": [issue]}
    db.commit()
    return batch


def _manifest(workbook) -> WorkbookManifest:
    if SYSTEM_SHEET not in workbook.sheetnames:
        raise ValueError("missing system sheet")
    raw = workbook[SYSTEM_SHEET]["A2"].value
    if not isinstance(raw, str):
        raise ValueError("missing manifest")
    payload = json.loads(raw)
    if payload.get("contract_version") != 1:
        raise UnsupportedContract
    return WorkbookManifest.model_validate(payload)


def _sheet_contract(manifest: WorkbookManifest, sheet_name: str):
    fields = manifest.product_fields if sheet_name == PRODUCT_SHEET else manifest.price_fields
    technical = TECH_PRODUCT_HEADERS if sheet_name == PRODUCT_SHEET else TECH_PRICE_HEADERS
    headers = [field.label for field in fields] + list(technical)
    keys = [field.key for field in fields] + [TECH_KEYS[label] for label in technical]
    return headers, dict(zip(headers, keys, strict=True))


def _header_issues(sheet_name: str, actual: list[str], expected_headers: list[str]) -> list[dict]:
    issues: list[dict] = []
    duplicates = sorted({value for value in actual if value and actual.count(value) > 1})
    if duplicates:
        issues.append(
            _issue(
                "DUPLICATE_HEADER",
                "表头重复：" + "、".join(duplicates),
                sheet=sheet_name,
                row=HEADER_ROW,
            )
        )
    missing = [value for value in expected_headers if value not in actual]
    if missing:
        issues.append(
            _issue(
                "MISSING_HEADER",
                "缺少表头：" + "、".join(missing),
                sheet=sheet_name,
                row=HEADER_ROW,
            )
        )
    unknown = [value for value in actual if value and value not in expected_headers]
    if unknown:
        issues.append(
            _issue(
                "UNKNOWN_HEADER",
                "存在未知表头：" + "、".join(unknown),
                sheet=sheet_name,
                row=HEADER_ROW,
            )
        )
    if not duplicates and not missing and not unknown and actual != expected_headers:
        issues.append(
            _issue(
                "HEADER_ORDER_CHANGED",
                "表头顺序已改变，请重新下载模板",
                sheet=sheet_name,
                row=HEADER_ROW,
            )
        )
    return issues


def parse_workbook(
    db: Session, file_bytes: bytes, filename: str, user_id: int
) -> ImportBatch:
    """Persist raw rows and safe structural issues; never execute workbook formulas."""

    if len(file_bytes) > MAX_WORKBOOK_BYTES:
        batch = _new_batch(db, file_bytes, filename, user_id)
        return _finish_failed(
            db,
            batch,
            _issue("FILE_TOO_LARGE", "文件超过 8 MB，请减少空白行或拆分后重试"),
            "Excel 文件超过 8 MB，未读取任何业务数据",
        )
    try:
        workbook = load_workbook(BytesIO(file_bytes), data_only=False, read_only=True)
    except Exception:
        batch = _new_batch(db, file_bytes, filename, user_id)
        return _finish_failed(
            db,
            batch,
            _issue("MALFORMED_WORKBOOK", "无法读取 Excel 工作簿，请重新下载模板后填写"),
            "无法读取 Excel 工作簿，请重新下载模板后填写",
        )
    try:
        manifest = _manifest(workbook)
    except UnsupportedContract:
        batch = _new_batch(db, file_bytes, filename, user_id)
        return _finish_failed(
            db,
            batch,
            _issue("STALE_CONTRACT", "模板版本已不受支持，请重新下载最新模板"),
            "模板版本已不受支持，请重新下载最新模板",
        )
    except (ValueError, json.JSONDecodeError, PydanticValidationError):
        batch = _new_batch(db, file_bytes, filename, user_id)
        return _finish_failed(
            db,
            batch,
            _issue("INVALID_MANIFEST", "模板系统信息缺失或已损坏，请重新下载模板"),
            "模板系统信息缺失或已损坏，请重新下载模板",
        )

    batch = _new_batch(db, file_bytes, filename, user_id, manifest=manifest)
    global_issues: list[dict] = []
    expected_sheets = {PRODUCT_SHEET, PRICE_SHEET, SYSTEM_SHEET}
    missing_sheets = [name for name in expected_sheets if name not in workbook.sheetnames]
    if missing_sheets:
        global_issues.append(
            _issue("MISSING_SHEET", "缺少工作表：" + "、".join(sorted(missing_sheets)))
        )
    unexpected = [name for name in workbook.sheetnames if name not in expected_sheets]
    if unexpected:
        global_issues.append(
            _issue("UNEXPECTED_SHEET", "存在未知工作表：" + "、".join(unexpected))
        )
    catalog = load_product_catalog(db)
    if manifest.product_schema_version != catalog.table.schema_version:
        global_issues.append(
            _issue(
                "STALE_SCHEMA",
                "产品字段结构已变化，请下载最新模板；当前文件不会提交",
            )
        )

    staged: list[ImportRow] = []
    for sheet_name, sheet_key in SHEETS.items():
        if sheet_name not in workbook.sheetnames:
            continue
        sheet = workbook[sheet_name]
        expected_headers, header_keys = _sheet_contract(manifest, sheet_name)
        rows_iter = sheet.iter_rows(min_row=HEADER_ROW, max_col=sheet.max_column)
        try:
            header_cells = next(rows_iter)
        except StopIteration:
            header_cells = ()
        actual_headers = [
            str(cell.value).strip() if cell.value is not None else "" for cell in header_cells
        ]
        global_issues.extend(_header_issues(sheet_name, actual_headers, expected_headers))
        for cells in rows_iter:
            source_row = cells[0].row if cells else DATA_START_ROW
            if all(cell.value in (None, "") for cell in cells):
                continue
            raw: dict[str, Any] = {}
            row_issues: list[dict] = []
            for header, cell in zip(actual_headers, cells, strict=True):
                key = header_keys.get(header)
                if key is None or key in raw:
                    continue
                raw[key] = _json_value(cell.value)
                if cell.data_type == "f":
                    row_issues.append(
                        _issue(
                            "FORMULA_NOT_ALLOWED",
                            "不允许使用公式，请粘贴计算后的值",
                            sheet=sheet_name,
                            row=source_row,
                            field=header,
                        )
                    )
            staged.append(
                ImportRow(
                    batch_id=batch.id,
                    sheet_key=sheet_key,
                    source_row_number=source_row,
                    raw_values=raw,
                    normalized_values={},
                    snapshot={},
                    action="block" if row_issues else "pending",
                    issues=row_issues,
                )
            )
    db.add_all(staged)
    batch.total_rows = len(staged)
    batch.result = {"issues": global_issues}
    batch.block_count = len(global_issues) + sum(bool(row.issues) for row in staged)
    if any(issue["code"] != "STALE_SCHEMA" for issue in global_issues):
        batch.status = "failed"
        batch.error_message = "工作簿结构不符合当前模板，请根据问题提示修正"
    db.commit()
    return batch
