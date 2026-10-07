"""Generate the protected, schema-versioned two-sheet product workbook."""

from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select
from sqlalchemy.orm import Session

from domains.dynamic_data import models as dynamic_models
from domains.dynamic_data import permissions as dynamic_permissions
from domains.dynamic_data.schemas import Actor
from domains.masterdata import models as master_models

from .catalog import PRICE_FIELDS, CoreField, ProductCatalog, load_product_catalog
from .contracts import WorkbookManifest

PRODUCT_SHEET = "产品资料"
PRICE_SHEET = "价格记录"
SYSTEM_SHEET = "_系统信息"
HEADER_ROW = 4
DATA_START_ROW = 5
INPUT_ROWS = 2000
TECH_PRODUCT_HEADERS = ("__产品ID", "__产品版本", "__扩展版本")
TECH_PRICE_HEADERS = ("__价格区间ID", "__价格区间版本")

HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
TITLE_FILL = PatternFill("solid", fgColor="0B5CAB")
GUIDE_FILL = PatternFill("solid", fgColor="EEF5FA")


def _choice_values(db: Session, model) -> list[str]:
    return sorted(
        {
            str(value).strip()
            for value in db.scalars(
                select(model.name).where(model.status.is_(True)).order_by(model.name)
            )
            if value and str(value).strip()
        }
    )


def _currencies(db: Session) -> list[str]:
    values: set[str] = set()
    for statement in (
        select(master_models.ExchangeRate.from_currency),
        select(master_models.ExchangeRate.to_currency),
        select(master_models.Product.currency),
        select(master_models.ProductPricePeriod.currency),
    ):
        values.update(
            str(value).strip().upper()
            for value in db.scalars(statement)
            if value and len(str(value).strip()) == 3
        )
    return sorted(values)


class ChoiceStore:
    def __init__(self, sheet):
        self.sheet = sheet
        self.column = 2

    def add(self, key: str, values: list[str]) -> str | None:
        values = list(dict.fromkeys(values))
        if not values:
            return None
        column = self.column
        self.column += 1
        self.sheet.cell(4, column, key)
        for row, value in enumerate(values, 5):
            self.sheet.cell(row, column, value)
        letter = get_column_letter(column)
        return f"'{SYSTEM_SHEET}'!${letter}$5:${letter}${4 + len(values)}"


def _setup_sheet(sheet, title: str, guidance: str, headers: list[str]) -> None:
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = sheet.cell(1, 1, title)
    title_cell.font = Font(color="FFFFFF", bold=True, size=14)
    title_cell.fill = TITLE_FILL
    title_cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 26
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    guide_cell = sheet.cell(2, 1, guidance)
    guide_cell.fill = GUIDE_FILL
    guide_cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[2].height = 34
    for column, header in enumerate(headers, 1):
        cell = sheet.cell(HEADER_ROW, column, header)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = f"A{DATA_START_ROW}"
    sheet.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(len(headers))}{HEADER_ROW}"
    sheet.protection.sheet = True
    sheet.protection.insertRows = False
    sheet.protection.deleteRows = False
    for column in range(1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 18
    sheet.column_dimensions["C"].width = 30


def _unlock_and_format(sheet, column_count: int, formats: dict[int, str]) -> None:
    for row in range(DATA_START_ROW, DATA_START_ROW + INPUT_ROWS):
        for column in range(1, column_count + 1):
            cell = sheet.cell(row, column)
            cell.protection = Protection(locked=False)
            if column in formats:
                cell.number_format = formats[column]


def _add_list_validation(sheet, column: int, formula: str | None) -> None:
    if formula is None:
        return
    validation = DataValidation(type="list", formula1=formula, allow_blank=True)
    validation.error = "请选择下拉列表中的有效值"
    validation.errorTitle = "选择无效"
    validation.showErrorMessage = True
    sheet.add_data_validation(validation)
    letter = get_column_letter(column)
    validation.add(f"{letter}{DATA_START_ROW}:{letter}{DATA_START_ROW + INPUT_ROWS - 1}")


def _add_price_validations(price_sheet, choices: ChoiceStore, currencies: list[str]) -> None:
    _add_list_validation(price_sheet, 3, choices.add("价格类型", ["采购价", "卖价"]))
    amount = DataValidation(
        type="decimal", operator="between", formula1="0", formula2="99999999.99", allow_blank=True
    )
    amount.error = "价格必须是 0 至 99999999.99 的数字"
    amount.showErrorMessage = True
    price_sheet.add_data_validation(amount)
    amount.add(f"D{DATA_START_ROW}:D{DATA_START_ROW + INPUT_ROWS - 1}")
    currency_range = choices.add("币种", currencies)
    if currency_range:
        _add_list_validation(price_sheet, 5, currency_range)
    else:
        currency = DataValidation(
            type="custom",
            formula1=(
                f'=OR(E{DATA_START_ROW}="",AND(LEN(E{DATA_START_ROW})=3,'
                f'EXACT(E{DATA_START_ROW},UPPER(E{DATA_START_ROW}))))'
            ),
            allow_blank=True,
        )
        currency.error = "币种请填写三个大写英文字母，例如 JPY"
        currency.showErrorMessage = True
        price_sheet.add_data_validation(currency)
        currency.add(f"E{DATA_START_ROW}:E{DATA_START_ROW + INPUT_ROWS - 1}")
    for column in (6, 7):
        validation = DataValidation(
            type="date",
            operator="between",
            formula1="DATE(1900,1,1)",
            formula2="DATE(9999,12,31)",
            allow_blank=True,
        )
        validation.error = "请填写有效日期，格式为 YYYY-MM-DD"
        validation.showErrorMessage = True
        price_sheet.add_data_validation(validation)
        letter = get_column_letter(column)
        validation.add(f"{letter}{DATA_START_ROW}:{letter}{DATA_START_ROW + INPUT_ROWS - 1}")


def _extension_value(field, value: Any) -> Any:
    if value is None:
        return None
    if field.field_type in {"single_select", "multi_select"}:
        labels = {str(option["id"]): option["label"] for option in field.config["options"]}
        if field.field_type == "single_select":
            return labels.get(str(value), str(value))
        return "；".join(labels.get(str(item), str(item)) for item in value)
    if field.field_type == "boolean":
        return "是" if value else "否"
    return value


def _product_rows(db: Session, catalog: ProductCatalog, product_ids: list[int] | None):
    statement = select(master_models.Product)
    if product_ids is not None:
        statement = statement.where(master_models.Product.id.in_(set(product_ids)))
    products = list(db.scalars(statement.order_by(master_models.Product.id)))
    if not products:
        return products, {}, {}, {}
    ids = [product.id for product in products]
    relations = {
        "port": dict(
            db.execute(
                select(master_models.Port.id, master_models.Port.name).where(
                    master_models.Port.id.in_({p.port_id for p in products if p.port_id})
                )
            ).all()
        ),
        "supplier": dict(
            db.execute(
                select(master_models.Supplier.id, master_models.Supplier.name).where(
                    master_models.Supplier.id.in_({p.supplier_id for p in products if p.supplier_id})
                )
            ).all()
        ),
        "category": dict(
            db.execute(
                select(master_models.Category.id, master_models.Category.name).where(
                    master_models.Category.id.in_({p.category_id for p in products if p.category_id})
                )
            ).all()
        ),
        "country": dict(
            db.execute(
                select(master_models.Country.id, master_models.Country.name).where(
                    master_models.Country.id.in_({p.country_id for p in products if p.country_id})
                )
            ).all()
        ),
    }
    anchors = {
        record.source_record_id: record
        for record in db.scalars(
            select(dynamic_models.DataRecord).where(
                dynamic_models.DataRecord.table_id == catalog.table.id,
                dynamic_models.DataRecord.source_record_id.in_(
                    [str(product_id) for product_id in ids]
                ),
                dynamic_models.DataRecord.status == "active",
            )
        )
        if record.source_record_id is not None
    }
    periods = list(
        db.scalars(
            select(master_models.ProductPricePeriod)
            .where(
                master_models.ProductPricePeriod.product_id.in_(ids),
                master_models.ProductPricePeriod.status.is_(True),
            )
            .order_by(
                master_models.ProductPricePeriod.product_id,
                master_models.ProductPricePeriod.effective_from,
                master_models.ProductPricePeriod.price_type,
                master_models.ProductPricePeriod.id,
            )
        )
    )
    return products, relations, anchors, periods


def _core_value(field: CoreField, product: master_models.Product, relations) -> Any:
    value = getattr(product, field.attribute)
    if field.relation:
        return relations[field.relation].get(value)
    if field.key == "status":
        return "启用" if value else "停用"
    return value


def build_product_workbook(
    db: Session,
    actor: Actor,
    *,
    include_existing: bool,
    product_ids: list[int] | None = None,
) -> bytes:
    dynamic_permissions.require_admin(actor)
    if not include_existing and product_ids:
        raise ValueError("空白模板不能指定已有产品")
    catalog = load_product_catalog(db)
    manifest = WorkbookManifest(
        product_schema_version=catalog.table.schema_version,
        generated_at=datetime.now(UTC),
        product_fields=catalog.product_contracts,
        price_fields=list(PRICE_FIELDS),
    )

    wb = Workbook()
    product_sheet = wb.active
    product_sheet.title = PRODUCT_SHEET
    price_sheet = wb.create_sheet(PRICE_SHEET)
    system_sheet = wb.create_sheet(SYSTEM_SHEET)
    system_sheet["A1"] = "workbook_manifest"
    system_sheet["A2"] = manifest.model_dump_json()
    choices = ChoiceStore(system_sheet)

    product_headers = [field.label for field in catalog.product_contracts] + list(
        TECH_PRODUCT_HEADERS
    )
    price_headers = [field.label for field in PRICE_FIELDS] + list(TECH_PRICE_HEADERS)
    _setup_sheet(
        product_sheet,
        "产品资料",
        "每个产品一行；产品代码与港口用于识别唯一产品。空白表示保留原值，清空请填写 __CLEAR__。",
        product_headers,
    )
    _setup_sheet(
        price_sheet,
        "价格记录",
        "每个价格区间一行；同一产品、同一价格类型的日期区间不能重叠。",
        price_headers,
    )
    _unlock_and_format(product_sheet, len(product_headers), {1: "@"})
    _unlock_and_format(price_sheet, len(price_headers), {1: "@", 4: "0.00", 6: "yyyy-mm-dd", 7: "yyyy-mm-dd"})

    master_choices = {
        "port": _choice_values(db, master_models.Port),
        "supplier": _choice_values(db, master_models.Supplier),
        "category": _choice_values(db, master_models.Category),
        "country": _choice_values(db, master_models.Country),
    }
    for column, contract in enumerate(catalog.product_contracts, 1):
        if contract.key in master_choices:
            _add_list_validation(
                product_sheet,
                column,
                choices.add(contract.label, master_choices[contract.key]),
            )
        elif contract.key == "status":
            _add_list_validation(
                product_sheet, column, choices.add("产品状态", ["启用", "停用"])
            )
        elif contract.key.startswith("extension:"):
            field = next(
                item for item in catalog.extension_fields if f"extension:{item.id}" == contract.key
            )
            if field.field_type in {"single_select", "multi_select"}:
                values = [
                    option["label"]
                    for option in field.config.get("options", [])
                    if option.get("active", True)
                ]
                _add_list_validation(
                    product_sheet, column, choices.add(f"扩展_{field.id}", values)
                )
            elif field.field_type == "boolean":
                _add_list_validation(
                    product_sheet, column, choices.add(f"布尔_{field.id}", ["是", "否"])
                )
            elif field.field_type == "date":
                for row in range(DATA_START_ROW, DATA_START_ROW + INPUT_ROWS):
                    product_sheet.cell(row, column).number_format = "yyyy-mm-dd"
            elif field.field_type == "datetime":
                for row in range(DATA_START_ROW, DATA_START_ROW + INPUT_ROWS):
                    product_sheet.cell(row, column).number_format = "yyyy-mm-dd hh:mm:ss"
            elif field.field_type == "number":
                for row in range(DATA_START_ROW, DATA_START_ROW + INPUT_ROWS):
                    product_sheet.cell(row, column).number_format = "0.######"
    _add_list_validation(price_sheet, 2, choices.add("价格港口", master_choices["port"]))
    _add_price_validations(price_sheet, choices, _currencies(db))

    products, relations, anchors, periods = (
        _product_rows(db, catalog, product_ids) if include_existing else ([], {}, {}, [])
    )
    prefix = list(catalog.core_prefix)
    suffix = list(catalog.core_suffix)
    for row_number, product in enumerate(products, DATA_START_ROW):
        anchor = anchors.get(str(product.id))
        values = [_core_value(field, product, relations) for field in prefix]
        values.extend(
            _extension_value(field, (anchor.values if anchor else {}).get(str(field.id)))
            for field in catalog.extension_fields
        )
        values.extend(_core_value(field, product, relations) for field in suffix)
        values.extend([product.id, product.revision, anchor.revision if anchor else 0])
        for column, value in enumerate(values, 1):
            product_sheet.cell(row_number, column, value)

    products_by_id = {product.id: product for product in products}
    for row_number, period in enumerate(periods, DATA_START_ROW):
        product = products_by_id[period.product_id]
        values = [
            product.code,
            relations["port"].get(product.port_id),
            "采购价" if period.price_type == "purchase" else "卖价",
            float(period.amount),
            period.currency,
            period.effective_from,
            period.effective_to,
            period.id,
            period.revision,
        ]
        for column, value in enumerate(values, 1):
            price_sheet.cell(row_number, column, value)

    for sheet, headers, hidden_count in (
        (product_sheet, product_headers, len(TECH_PRODUCT_HEADERS)),
        (price_sheet, price_headers, len(TECH_PRICE_HEADERS)),
    ):
        for column in range(len(headers) - hidden_count + 1, len(headers) + 1):
            sheet.column_dimensions[get_column_letter(column)].hidden = True
            for row in range(DATA_START_ROW, DATA_START_ROW + INPUT_ROWS):
                sheet.cell(row, column).protection = Protection(locked=True)

    system_sheet.sheet_state = "veryHidden"
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
