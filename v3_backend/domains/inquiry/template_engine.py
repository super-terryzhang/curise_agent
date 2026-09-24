"""Excel rendering engine — turns a SupplierTemplate + order data into bytes.

Two paths:
1. Template-driven (`field_positions` + `product_table_config` on SupplierTemplate):
   loads the original .xlsx file (preserves formatting/formulas) and fills cells.
2. Generic fallback: builds a plain inquiry sheet with title / metadata / product
   table — used when no template is bound to the supplier.

Phase 5 ships the deterministic engine. Phase 6 may layer an agentic
cell-by-cell writer on top (v2's InquiryWorkbook); the public surface here
returns Excel bytes synchronously, callers don't need to know which path ran.
"""

from __future__ import annotations

import contextlib
import io
import logging
import re
from copy import copy
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

from domains.inquiry.models import SupplierTemplate
from domains.inquiry.template_contract import (
    assert_source_workbook_matches,
    canonical_metadata_key,
    normalized_contract,
)

logger = logging.getLogger(__name__)

_WARNING_FILL = PatternFill(fill_type="solid", fgColor="FFF2CC")


def render_inquiry_excel(
    template: SupplierTemplate | None,
    *,
    metadata: dict[str, Any],
    products: list[dict[str, Any]],
    supplier_id: int,
    template_file_bytes: bytes | None = None,
    field_mapping: dict[str, str] | None = None,
) -> bytes:
    """Render the inquiry workbook for one supplier.

    - When a `template` with `field_positions` is provided AND we have the
      original .xlsx bytes, load + fill (preserves styles/formulas).
    - When `template` exists but there's no source file, build a fresh workbook
      using the template config.
    - Otherwise fall back to the generic layout.
    """
    if template is not None and template.field_positions and template_file_bytes:
        wb = load_workbook(io.BytesIO(template_file_bytes))
        ws = wb.active
        contract = normalized_contract(template)
        if contract["is_dynamic"]:
            assert_source_workbook_matches(ws, template)
            _fill_zoned_template(ws, template, metadata, products, field_mapping)
        else:
            _fill_with_template(ws, template, metadata, products, field_mapping)
    elif template is not None and template.field_positions:
        wb = Workbook()
        ws = wb.active
        ws.title = "Inquiry"
        _fill_with_template(ws, template, metadata, products, field_mapping)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Inquiry"
        _fill_generic(ws, metadata, products, supplier_id)

    automated = bool(products) and all("rfq_quantity" in p for p in products)
    arrangement_sources = any(p.get("source_po_number") for p in products)
    if automated:
        from domains.inquiry.workbook_quality import finalize_layout, save_with_formula_cache
        if template is not None:
            finalize_layout(ws, template, products)
        # Automated single-PO and arrangement workbooks both require the
        # source sheet; workbook_quality validates its presence.
        _append_source_sheet(wb, metadata, products)
        return save_with_formula_cache(wb)
    if arrangement_sources:
        _append_source_sheet(wb, metadata, products)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ─── Template-driven fill ─────────────────────────────────────


def _fill_zoned_template(
    ws: Any,
    template: SupplierTemplate,
    metadata: dict[str, Any],
    products: list[dict[str, Any]],
    field_mapping: dict[str, str] | None,
) -> None:
    """Fill a contract-backed workbook and resize its product/summary zones.

    The template is treated as layout, not data.  Header samples are cleared,
    product rows are resized to the actual item count, and every formula is
    emitted from the declarative contract instead of copied opportunistically
    from whichever sample rows happen to contain formulas.
    """

    contract = normalized_contract(template)
    styles = contract["styles"]
    product_start = int(contract["product_start"])
    product_end = int(contract["product_end"])
    summary_start = int(contract["summary_start"])
    columns: dict[str, str] = contract["columns"]
    product_formulas: dict[str, str] = contract["product_row_formulas"]

    _write_header_fields(ws, template, metadata, field_mapping)
    for position, value in (styles.get("static_values") or {}).items():
        cell = ws[str(position).upper()]
        if not isinstance(cell, MergedCell):
            cell.value = value
            if isinstance(value, str):
                cell.data_type = "s"

    output_rows = max(len(products), 1)
    original_capacity = product_end - product_start + 1
    row_delta = output_rows - original_capacity
    new_product_end = product_start + output_rows - 1
    _resize_product_zone(
        ws,
        product_start=product_start,
        product_end=product_end,
        summary_start=summary_start,
        output_rows=output_rows,
        product_row_merges=styles.get("product_row_merges") or [],
    )

    writable_columns = set(columns) | set(product_formulas)
    for row in range(product_start, new_product_end + 1):
        for col in writable_columns:
            cell = ws[f"{col}{row}"]
            if not isinstance(cell, MergedCell):
                cell.value = None

    for index, product in enumerate(products):
        row = product_start + index
        matched = product.get("matched_product") or {}
        for col, field_key in columns.items():
            value = _resolve_product_value(field_key, index, product, matched, metadata)
            if value in (None, ""):
                continue
            if field_key in ("quantity", "unit_price", "total_price", "amount"):
                value = _to_number_or_none(value)
                if value is None:
                    continue
            cell = ws[f"{col}{row}"]
            if isinstance(cell, MergedCell):
                raise ValueError(f"TEMPLATE_DATA_CELL_MERGED: {col}{row}")
            cell.value = value
            if isinstance(value, str):
                cell.data_type = "s"
            if field_key in ("unit_price", "total_price", "amount"):
                cell.number_format = "#,##0.00"
        for col, formula_template in product_formulas.items():
            cell = ws[f"{col}{row}"]
            if isinstance(cell, MergedCell):
                raise ValueError(f"TEMPLATE_FORMULA_CELL_MERGED: {col}{row}")
            cell.value = formula_template.replace("{row}", str(row))
        _highlight_warning_row(ws, row, product, writable_columns)

    # Keep one structurally valid blank product row when no products exist.
    if not products:
        for col, formula_template in product_formulas.items():
            ws[f"{col}{product_start}"] = formula_template.replace(
                "{row}", str(product_start)
            )

    new_summary_start = summary_start + row_delta
    new_summary_end = int(contract["summary_end"]) + row_delta
    for row in range(new_summary_start, new_summary_end + 1):
        for col in styles.get("stale_columns_in_summary") or []:
            cell = ws[f"{str(col).upper()}{row}"]
            if not isinstance(cell, MergedCell):
                cell.value = None

    for position, value in (styles.get("summary_static_values") or {}).items():
        new_position = _shift_cell_ref(str(position), summary_start, row_delta)
        cell = ws[new_position]
        if not isinstance(cell, MergedCell):
            cell.value = value
            if isinstance(value, str):
                cell.data_type = "s"

    formula_refs: dict[str, str] = {}
    for item in styles.get("summary_formulas") or []:
        source_position = str(item.get("cell") or "").upper()
        if not source_position:
            continue
        target_position = _shift_cell_ref(source_position, summary_start, row_delta)
        col = str(item.get("col") or _cell_column(target_position)).upper()
        if item.get("type") == "product_sum":
            formula = f"=SUM({col}{product_start}:{col}{new_product_end})"
            formula_refs["sum_cell"] = target_position
        elif item.get("type") == "relative":
            formula = str(item.get("formula_template") or "")
            for key, ref in formula_refs.items():
                formula = formula.replace(f"{{{key}}}", ref)
        else:
            raise ValueError(f"TEMPLATE_SUMMARY_FORMULA_UNSUPPORTED: {item.get('type')}")
        if "{" in formula or not formula.startswith("="):
            raise ValueError(f"TEMPLATE_SUMMARY_FORMULA_INCOMPLETE: {target_position}")
        ws[target_position] = formula
        label = str(item.get("label") or "").lower()
        if "tax" in label:
            formula_refs["tax_cell"] = target_position
        if "grand" in label or ("total" in label and "sub" not in label):
            formula_refs["grand_total_cell"] = target_position

    for item in styles.get("external_refs") or []:
        position = str(item.get("cell") or "").upper()
        formula = str(item.get("formula_template") or "")
        for key, ref in formula_refs.items():
            formula = formula.replace(f"{{{key}}}", ref)
        if not position or "{" in formula or not formula.startswith("="):
            raise ValueError("TEMPLATE_EXTERNAL_REFERENCE_INCOMPLETE")
        ws[position] = formula


def _write_header_fields(
    ws: Any,
    template: SupplierTemplate,
    metadata: dict[str, Any],
    field_mapping: dict[str, str] | None,
) -> None:
    styles = template.template_styles or {}
    configured = styles.get("header_fields")
    formats = styles.get("header_formats") or {}
    specs: list[tuple[str, str, str | None]] = []
    if isinstance(configured, dict) and configured:
        for position, path in configured.items():
            specs.append(
                (
                    str(position).upper(),
                    canonical_metadata_key(str(path)),
                    formats.get(position) or formats.get(str(position).upper()),
                )
            )
    else:
        for key, raw in (template.field_positions or {}).items():
            info = raw if isinstance(raw, dict) else {"position": raw}
            position = str(info.get("position") or "").upper()
            mapped = info.get("source") or (field_mapping or {}).get(key, key)
            specs.append((position, canonical_metadata_key(str(mapped)), info.get("format")))

    for position, _, _ in specs:
        if not position:
            continue
        cell = ws[position]
        if isinstance(cell, MergedCell):
            raise ValueError(f"TEMPLATE_HEADER_CELL_MERGED_CHILD: {position}")
        cell.value = None
    for position, source, display_format in specs:
        value = metadata.get(source)
        if value in (None, ""):
            continue
        if display_format:
            value = str(display_format).replace("{value}", str(value))
        cell = ws[position]
        cell.value = value
        if isinstance(value, str):
            cell.data_type = "s"


def _resize_product_zone(
    ws: Any,
    *,
    product_start: int,
    product_end: int,
    summary_start: int,
    output_rows: int,
    product_row_merges: list[dict[str, Any]],
) -> None:
    original_max_row = ws.max_row
    original_heights = {
        row: dim.height
        for row, dim in ws.row_dimensions.items()
        if dim.height is not None
    }
    movable_merges: list[tuple[int, int, int, int]] = []
    for merged in list(ws.merged_cells.ranges):
        min_col, min_row, max_col, max_row = range_boundaries(str(merged))
        if max_row < product_start:
            continue
        if min_row < product_start <= max_row:
            raise ValueError(f"TEMPLATE_MERGE_CROSSES_PRODUCT_ZONE: {merged}")
        ws.unmerge_cells(str(merged))
        movable_merges.append((min_col, min_row, max_col, max_row))

    capacity = product_end - product_start + 1
    row_delta = output_rows - capacity
    if row_delta > 0:
        ws.insert_rows(summary_start, row_delta)
        for row in range(summary_start, summary_start + row_delta):
            _copy_row_style(ws, product_start, row)
    elif row_delta < 0:
        ws.delete_rows(product_start + output_rows, -row_delta)

    for row in list(ws.row_dimensions):
        if row >= product_start:
            del ws.row_dimensions[row]
    product_height = original_heights.get(product_start)
    if product_height is not None:
        for row in range(product_start, product_start + output_rows):
            ws.row_dimensions[row].height = product_height
    for original_row, height in original_heights.items():
        if original_row >= summary_start:
            ws.row_dimensions[original_row + row_delta].height = height

    for min_col, min_row, max_col, max_row in movable_merges:
        if product_start <= min_row and max_row <= product_end:
            continue
        if min_row < summary_start:
            raise ValueError(
                "TEMPLATE_MERGE_BETWEEN_ZONES: "
                f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"
            )
        _merge_if_needed(
            ws,
            min_row + row_delta,
            min_col,
            max_row + row_delta,
            max_col,
        )

    for row in range(product_start, product_start + output_rows):
        for item in product_row_merges:
            _merge_if_needed(
                ws,
                row,
                int(item["start_col"]),
                row,
                int(item["end_col"]),
            )

    # Clear obsolete dimensions that row deletion can leave behind.
    for row in range(original_max_row + row_delta + 1, original_max_row + 2):
        if row in ws.row_dimensions:
            del ws.row_dimensions[row]


def _copy_row_style(ws: Any, source_row: int, target_row: int) -> None:
    for col in range(1, ws.max_column + 1):
        source = ws.cell(source_row, col)
        target = ws.cell(target_row, col)
        if source.has_style:
            target._style = copy(source._style)  # noqa: SLF001 - openpyxl style clone
        if source.number_format:
            target.number_format = source.number_format


def _merge_if_needed(
    ws: Any, min_row: int, min_col: int, max_row: int, max_col: int
) -> None:
    if min_row == max_row and min_col == max_col:
        return
    ref = (
        f"{get_column_letter(min_col)}{min_row}:"
        f"{get_column_letter(max_col)}{max_row}"
    )
    if ref not in {str(item) for item in ws.merged_cells.ranges}:
        ws.merge_cells(ref)


def _shift_cell_ref(position: str, threshold_row: int, delta: int) -> str:
    match = re.fullmatch(r"([A-Z]{1,3})(\d+)", position.upper())
    if not match:
        raise ValueError(f"TEMPLATE_CELL_REFERENCE_INVALID: {position}")
    col, row_text = match.groups()
    row = int(row_text)
    return f"{col}{row + delta if row >= threshold_row else row}"


def _cell_column(position: str) -> str:
    match = re.match(r"([A-Z]{1,3})", position.upper())
    if not match:
        raise ValueError(f"TEMPLATE_CELL_REFERENCE_INVALID: {position}")
    return match.group(1)


def _fill_with_template(
    ws: Any,
    template: SupplierTemplate,
    metadata: dict[str, Any],
    products: list[dict[str, Any]],
    field_mapping: dict[str, str] | None,
) -> None:
    field_positions = template.field_positions or {}
    table_config = template.product_table_config or {}
    formula_columns = {
        c.upper() for c in (table_config.get("formula_columns") or []) if isinstance(c, str)
    }

    # Header fields.
    # Two-pass: (1) force-blank every declared cell so stale sample data
    # baked into the supplier-uploaded template doesn't leak through, then
    # (2) write the non-empty values from metadata. Without the blank pass,
    # a template uploaded with a real address pre-filled in I8 will keep
    # showing that address on every new inquiry — see prod bug 2026-05-22
    # where order 106 (Osaka) showed "東京都中央区晴海" because the supplier
    # uploaded the template after rendering a Tokyo order and didn't clear
    # the cell.
    for field_key, pos_info in field_positions.items():
        position = pos_info if isinstance(pos_info, str) else pos_info.get("position", "")
        if not position:
            continue
        cell = ws[position]
        if isinstance(cell, MergedCell):
            logger.debug("Skipping merged cell %s for field %s", position, field_key)
            continue
        cell.value = None

    for field_key, pos_info in field_positions.items():
        position = pos_info if isinstance(pos_info, str) else pos_info.get("position", "")
        if not position:
            continue
        mapped_key = (field_mapping or {}).get(field_key, field_key)
        value = metadata.get(mapped_key, "")
        if value in (None, ""):
            continue
        cell = ws[position]
        if isinstance(cell, MergedCell):
            continue
        cell.value = value

    # Product rows
    start_row = int(table_config.get("start_row", 12))
    columns: dict[str, str] = table_config.get("columns") or {}

    # Suppliers in production sometimes upload templates with *example* product
    # rows already filled in (left over from a real order they cleaned up
    # before saving). If we don't strip those, every inquiry inherits a
    # cocktail of historical line items, e.g. supplier #1 with 1 real product
    # ends up shipping a sheet showing 83 rows of "確認中" / random fruit and a
    # #VALUE! grand total (because the leftover "確認中" string × price triggers
    # Excel's type error and the SUM range propagates it).
    last_dirty_row = _detect_filled_data_rows(ws, start_row, columns, formula_columns)
    if last_dirty_row >= start_row:
        _clear_data_cells(ws, start_row, last_dirty_row, columns, formula_columns)

    for i, product in enumerate(products):
        row_idx = start_row + i
        matched = product.get("matched_product") or {}
        for col_letter, field_key in columns.items():
            if col_letter.upper() in formula_columns:
                continue
            value = _resolve_product_value(field_key, i, product, matched, metadata)
            if value in (None, ""):
                continue
            # Quantity / unit_price / total_price must reach Excel as numbers
            # or the per-row formulas (=H*K) and SUM range explode into
            # #VALUE!. PDFs sometimes carry placeholder strings here (e.g.
            # "確認中" — Japanese for "TBD"). Coerce when possible, drop
            # otherwise so the cell stays blank (Excel treats blank as 0).
            if field_key in ("quantity", "unit_price", "total_price", "amount"):
                value = _to_number_or_none(value)
                if value is None:
                    continue
            cell = ws[f"{col_letter}{row_idx}"]
            if isinstance(cell, MergedCell):
                continue
            cell.value = value
            if isinstance(value, str):
                cell.data_type = "s"  # PO text must never become an Excel formula.
            if field_key in ("unit_price", "total_price", "amount"):
                cell.number_format = "0.00"
        _highlight_warning_row(ws, row_idx, product, set(columns) | formula_columns)


def _detect_filled_data_rows(
    ws: Any,
    start_row: int,
    columns: dict[str, str],
    formula_columns: set[str],
    max_scan: int = 500,
) -> int:
    """Return the last row in `start_row..start_row+max_scan` that has any
    non-formula data cell filled. Returns `start_row - 1` if nothing is
    filled — caller uses that to skip the clear pass.

    We ignore formula columns (those are part of the template's calculation
    machinery, not sample data) and merged cells (their value lives on the
    top-left anchor, scanning would misreport them).
    """
    last_filled = start_row - 1
    scan_cols = [c for c in columns if c.upper() not in formula_columns]
    for r in range(start_row, start_row + max_scan):
        # Subtotals are template content, not another example product row.
        if any(str(ws[f"{col}{r}"].value or "").strip().upper() in
               {"SUB TOTAL", "SUBTOTAL", "GRAND TOTAL", "小計", "合計"} for col in columns):
            break
        any_filled = False
        for col_letter in scan_cols:
            cell = ws[f"{col_letter}{r}"]
            if isinstance(cell, MergedCell):
                continue
            if cell.value not in (None, ""):
                any_filled = True
                break
        if any_filled:
            last_filled = r
        else:
            # First fully-empty row → assume the sample range ended here.
            break
    return last_filled


def _clear_data_cells(
    ws: Any,
    start_row: int,
    end_row: int,
    columns: dict[str, str],
    formula_columns: set[str],
) -> None:
    """Blank out all non-formula data cells in `start_row..end_row`.

    Formula columns are preserved — they're the per-row `=H*K` cells the
    template ships and we want them intact so totals continue to recalc.
    """
    for r in range(start_row, end_row + 1):
        for col_letter in columns:
            if col_letter.upper() in formula_columns:
                continue
            cell = ws[f"{col_letter}{r}"]
            if isinstance(cell, MergedCell):
                continue
            cell.value = None


def _to_number_or_none(value: Any) -> float | int | None:
    """Coerce a cell-bound value to a number, or `None` if it can't be one.

    Strings like "100", "100.5", " 100 " convert cleanly. Strings like
    "確認中" / "TBD" / "" do not — those become `None` and the caller skips
    the write so the cell stays blank. Numbers pass through untouched.
    """
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, str):
        s = value.strip()
        if not s:
            return None
        try:
            # `int()` would reject "100.5"; float() accepts both and Excel
            # is happy with floats in either case.
            return float(s)
        except ValueError:
            return None
    return None


def _resolve_product_value(
    field_key: str,
    index: int,
    product: dict[str, Any],
    matched: dict[str, Any],
    metadata: dict[str, Any],
) -> Any:
    """Single source of truth for filling product-row cells.

    Resolution rules (v44+):
      - `unit` / `pack_size` / `unit_price` → DB (matched_product) authoritative.
        For unit_price specifically: DB-ONLY, no fallback to Excel. The
        inquiry sheet ships to suppliers and must quote OUR catalog price,
        not the customer's PO price (prod incident 2026-05-20: customer
        complained inquiry showed 745 from their PO instead of 690 from DB).
      - quantity / line-item fields without a DB analog → extracted Excel
        values (no other source exists).
    """
    if field_key == "line_number":
        return index + 1
    if field_key == "quantity" and "rfq_quantity" in product:
        return product["rfq_quantity"]
    if field_key == "po_number":
        return (
            product.get("source_po_number")
            or product.get("po_number")
            or metadata.get("po_number", "")
        )
    if field_key == "currency":
        return product.get("currency") or matched.get("currency") or metadata.get("currency", "")
    if field_key == "description":
        return matched.get("pack_size") or product.get("description", "")
    if field_key == "product_name_en":
        return product.get("product_name") or matched.get("product_name_en", "")
    if field_key == "unit":
        # matched_product.unit is canonical (clean DB value);
        # extracted product.unit can be corrupted (e.g. "KG2.2").
        return product.get("rfq_unit") or matched.get("unit") or product.get("unit", "")
    if field_key == "unit_price":
        # v44 invariant: inquiry → suppliers MUST quote against our catalog
        # price (DB). Customer PO prices stay in `product.unit_price` for
        # anomaly detection / audit but never reach the inquiry sheet.
        # Treat 0 as a valid catalog price (free sample); only NULL falls
        # through to empty. NO fallback to Excel — explicit design choice.
        price = matched.get("price")
        return price if price is not None else ""
    val = product.get(field_key)
    return val if val is not None else matched.get(field_key, "")


def _append_source_sheet(wb, metadata, products):
    """Compact source cards keep demand and conversion readable in Excel viewers."""
    from decimal import Decimal

    name = "PO Source"
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    for row, text in [(1, "PO source and supplier quantities"),
                      (2, str(metadata.get("po_number", ""))),
                      (3, "Original demand is preserved; supplier quantities use the recorded basis.")]:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        cell = ws.cell(row, 1, text)
        cell.data_type = "s"
        cell.font = Font(size=14 if row == 1 else 11, bold=row < 3)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[row].height = 30
    for i, product in enumerate(products):
        start = 5+i*6
        proof = product.get("conversion_evidence") or {}
        base = ""
        supplier_quantity = product.get("rfq_quantity", product.get("quantity"))
        supplier_unit = product.get("rfq_unit") or product.get("unit")
        if proof.get("base_per_supplier_unit") and supplier_quantity is not None:
            base = f'{Decimal(str(supplier_quantity)) * Decimal(str(proof["base_per_supplier_unit"]))} {proof.get("base_unit", "")}'
        rows = [
            ["Source PO", product.get("source_po_number") or product.get("po_number"), "Product code", product.get("product_code")],
            ["PO line / page", f'{product.get("source_line", "")} / {product.get("page", "")}', "Source line ID", product.get("source_line_id")],
            ["Requested qty", _to_number_or_none(product.get("source_quantity", product.get("quantity"))), "Original unit", product.get("source_unit") or product.get("unit")],
            ["Supplier qty", supplier_quantity, "Supplier unit", supplier_unit],
            ["Base demand", base, "Catalog pack", (product.get("matched_product") or {}).get("pack_size")],
        ]
        for offset, values in enumerate(rows):
            for col, value in enumerate(values, 1):
                cell = ws.cell(start+offset, col, value)
                if isinstance(value, str):
                    cell.data_type = "s"
                cell.font = Font(size=11, bold=col in (1,3), color="FFFFFF" if col in (1,3) else "182B3A")
                if col in (1,3):
                    cell.fill = PatternFill("solid", fgColor="24445C")
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.row_dimensions[start+offset].height = 28
        ws.cell(start+5, 1, "Basis").font = Font(bold=True)
        ws.merge_cells(start_row=start+5,start_column=2,end_row=start+5,end_column=4)
        cell = ws.cell(start+5,2,proof.get("evidence", ""))
        cell.data_type = "s"
        cell.font = Font(size=10)
        cell.alignment = Alignment(wrap_text=True,vertical="top")
        ws.row_dimensions[start+5].height = 110
    for col, width in zip("ABCD", [22,24,22,32], strict=True):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A5"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.print_title_rows = "1:3"
    ws.print_area = f"A1:D{4+len(products)*6}"


# ─── Generic fallback layout ──────────────────────────────────


def _fill_generic(
    ws: Any, metadata: dict[str, Any], products: list[dict[str, Any]], supplier_id: int
) -> None:
    header_font = Font(bold=True, size=14)
    label_font = Font(bold=True, size=10)
    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    header_fill = PatternFill(start_color="D4A853", end_color="D4A853", fill_type="solid")
    header_text_font = Font(bold=True, size=10, color="FFFFFF")

    ws.merge_cells("A1:G1")
    ws["A1"] = "Purchase Order / 注文書"
    ws["A1"].font = header_font
    ws["A1"].alignment = Alignment(horizontal="center")

    meta_rows = [
        ("A3", "PO Number:", "B3", metadata.get("po_number", "")),
        ("A4", "Order Date:", "B4", metadata.get("order_date", "")),
        ("A5", "Delivery Date:", "B5", metadata.get("delivery_date", "")),
        ("A6", "Ship Name:", "B6", metadata.get("ship_name", "")),
        ("A7", "Currency:", "B7", metadata.get("currency", "")),
        ("D3", "Supplier ID:", "E3", str(supplier_id)),
        ("D4", "Vendor:", "E4", metadata.get("vendor_name", "")),
        ("D5", "Port:", "E5", metadata.get("destination_port", "")),
    ]
    for label_cell, label, value_cell, value in meta_rows:
        ws[label_cell] = label
        ws[label_cell].font = label_font
        ws[value_cell] = value

    table_start = 9
    headers = ["No.", "Product Code", "Product Name", "Qty", "Unit", "Unit Price", "Total"]
    col_widths = [6, 15, 35, 10, 8, 12, 12]
    for col_idx, (header, width) in enumerate(zip(headers, col_widths, strict=False), 1):
        cell = ws.cell(row=table_start, column=col_idx, value=header)
        cell.font = header_text_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = width

    total_amount = 0.0
    for i, product in enumerate(products, 1):
        row = table_start + i
        matched = product.get("matched_product") or {}
        # v44 invariant: DB-only, no Excel fallback. Customer PO prices in
        # `product.unit_price` are NEVER sent to suppliers (see
        # `_resolve_product_value` for the full rationale).
        db_price = matched.get("price")
        unit_price = db_price if db_price is not None else ""
        qty = product.get("rfq_quantity", product.get("quantity", ""))
        # Line total uses the DB price too. If DB price is missing this row
        # contributes nothing to the supplier inquiry total — the row will
        # show empty cells, signalling the catalog is incomplete for this
        # product and the buyer needs to fix it before sending.
        total_price: Any
        if db_price is None:
            total_price = ""
        else:
            try:
                total_price = float(qty) * float(db_price) if qty else ""
            except (TypeError, ValueError):
                total_price = ""
        with contextlib.suppress(TypeError, ValueError):
            total_amount += float(total_price)

        values = [
            i,
            product.get("product_code") or matched.get("code", ""),
            product.get("product_name") or matched.get("product_name_en", ""),
            qty,
            product.get("rfq_unit") or matched.get("unit") or product.get("unit", ""),
            unit_price,
            total_price,
        ]
        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(row=row, column=col_idx, value=value)
            cell.border = border
        _highlight_warning_row(ws, row, product, set(range(1, 8)))

    total_row = table_start + len(products) + 1
    ws.cell(row=total_row, column=5, value="Total:").font = label_font
    ws.cell(row=total_row, column=7, value=total_amount).font = label_font


def _highlight_warning_row(
    ws: Any,
    row: int,
    product: dict[str, Any],
    columns: set[str] | set[int],
) -> None:
    """Mark only included warning rows; excluded rows never reach the renderer."""
    if not product.get("inquiry_warnings") or not columns:
        return
    numeric_columns = sorted(
        column if isinstance(column, int) else ws[f"{column}{row}"].column
        for column in columns
    )
    for column in range(numeric_columns[0], numeric_columns[-1] + 1):
        cell = ws.cell(row=row, column=column)
        if not isinstance(cell, MergedCell):
            cell.fill = copy(_WARNING_FILL)


__all__ = ["render_inquiry_excel"]
