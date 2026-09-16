"""Bounded finalization of automated template output; never executes workbook code."""

import ast
import io
import re
import zipfile
from copy import copy
from decimal import Decimal
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils.cell import range_boundaries

from domains.inquiry.template_contract import normalized_contract


def finalize_layout(ws, template, products):
    config = normalized_contract(template)
    start = int(config["product_start"])
    last = start + len(products) - 1
    footers = [
        cell.row
        for row in ws.iter_rows(min_row=last + 1)
        for cell in row
        if cell.data_type == "f" and str(cell.value).upper().startswith("=SUM(")
    ]
    if footers:
        for row in range(last + 1, min(footers)):
            ws.row_dimensions[row].hidden = True
    for row in ws:
        for cell in row:
            if cell.data_type == "f":
                cell.number_format = "#,##0.00"
    for field in ("ship_name",):
        position = (template.field_positions or {}).get(field)
        if position:
            position = position if isinstance(position, str) else position["position"]
            cell = ws[position]
            font = copy(cell.font)
            font.sz = min(font.sz or 14, 18)
            cell.font = font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.row_dimensions[cell.row].height = max(ws.row_dimensions[cell.row].height or 20, 38)
    # Legacy templates without an explicit configured title remain RFQ output.
    # Contract-backed templates own their document title; do not rewrite it.
    explicit_title = (
        (template.template_styles or {}).get("static_values") or {}
    ).get("A1")
    if not explicit_title:
        for row in ws.iter_rows(max_row=min(start - 1, 16)):
            for cell in row:
                if (
                    isinstance(cell.value, str)
                    and cell.value.strip().upper() == "PURCHASE ORDER"
                ):
                    cell.value = "REQUEST FOR QUOTATION / 御見積依頼"
                    font = copy(cell.font)
                    font.sz = 20
                    cell.font = font
                    cell.alignment = Alignment(
                        horizontal="center", vertical="center", wrap_text=True
                    )
                    ws.row_dimensions[cell.row].height = 36
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.scale = None
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.print_title_rows = f"1:{start - 1}"
    ws.print_area = ws.calculate_dimension()


def formula_values(ws):
    """Evaluate only numeric references, arithmetic and SUM ranges; reject anything else.

    Keep formulas in the XLSX, filling their cached values for viewers that do
    not recalculate. Unsupported formulas fail automated output instead of
    displaying a stale or fabricated result.
    """
    cache, active = {}, set()

    def value(address):
        address = address.replace("$", "")
        if address in cache:
            return cache[address]
        if address in active or len(active) > 100:
            raise ValueError("TEMPLATE_FORMULA_CYCLE")
        cell = ws[address]
        if cell.data_type != "f":
            if cell.value in (None, ""):
                return Decimal(0)
            if isinstance(cell.value, bool) or not isinstance(cell.value, (int, float, Decimal)):
                raise ValueError("TEMPLATE_FORMULA_NON_NUMERIC_INPUT")
            result = Decimal(str(cell.value))
        else:
            active.add(address)
            expression = cell.value[1:].strip().upper()

            def sum_range(match):
                minimum_col, minimum_row, maximum_col, maximum_row = range_boundaries(match[1])
                if (maximum_col - minimum_col + 1) * (maximum_row - minimum_row + 1) > 10000:
                    raise ValueError("TEMPLATE_FORMULA_RANGE_LIMIT")
                return str(
                    sum(
                        (
                            value(c.coordinate)
                            for row in ws.iter_rows(
                                min_col=minimum_col,
                                max_col=maximum_col,
                                min_row=minimum_row,
                                max_row=maximum_row,
                            )
                            for c in row
                        ),
                        Decimal(0),
                    )
                )

            expression = re.sub(r"SUM\((\$?[A-Z]+\$?\d+:\$?[A-Z]+\$?\d+)\)", sum_range, expression)
            expression = re.sub(r"\$?[A-Z]{1,3}\$?\d+", lambda m: str(value(m[0])), expression)

            def arithmetic(node):
                if isinstance(node, ast.Constant) and type(node.value) in (int, float):
                    return Decimal(str(node.value))
                if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.USub, ast.UAdd)):
                    return (-1 if isinstance(node.op, ast.USub) else 1) * arithmetic(node.operand)
                if isinstance(node, ast.BinOp):
                    left, right = arithmetic(node.left), arithmetic(node.right)
                    if isinstance(node.op, ast.Add):
                        return left + right
                    if isinstance(node.op, ast.Sub):
                        return left - right
                    if isinstance(node.op, ast.Mult):
                        return left * right
                    if isinstance(node.op, ast.Div):
                        return left / right
                raise ValueError("TEMPLATE_FORMULA_UNSUPPORTED")

            if len(expression) > 2000:
                raise ValueError("TEMPLATE_FORMULA_LIMIT")
            result = arithmetic(ast.parse(expression, mode="eval").body)
            active.remove(address)
        if not result.is_finite():
            raise ValueError("TEMPLATE_FORMULA_NON_FINITE")
        cache[address] = result
        return result

    formulas = [cell for row in ws for cell in row if cell.data_type == "f"]
    if len(formulas) > 5000:
        raise ValueError("TEMPLATE_FORMULA_LIMIT")
    return {cell.coordinate: str(value(cell.coordinate)) for cell in formulas}


def save_with_formula_cache(wb):
    values = [formula_values(ws) for ws in wb.worksheets]
    raw = io.BytesIO()
    wb.save(raw)
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    # ElementTree otherwise rewrites the workbook namespace as ``ns0``. While
    # valid XML, macOS Quick Look and some older spreadsheet viewers reject
    # that form. Preserve the conventional default namespace used by Excel.
    ET.register_namespace("", ns)
    output = io.BytesIO()
    with (
        zipfile.ZipFile(raw) as source,
        zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for info in source.infolist():
            content = source.read(info.filename)
            match = re.fullmatch(r"xl/worksheets/sheet(\d+)\.xml", info.filename)
            if match:
                cells = values[int(match[1]) - 1]
                if cells:
                    root = ET.fromstring(content)
                    for cell in root.iter(f"{{{ns}}}c"):
                        if cell.get("r") in cells:
                            cached = cell.find(f"{{{ns}}}v")
                            if cached is None:
                                cached = ET.SubElement(cell, f"{{{ns}}}v")
                            cached.text = cells[cell.get("r")]
                    content = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(info, content)
    return output.getvalue()


def validate_output(content, template, products, metadata):
    """Read the saved workbook back; a skipped/merged data cell is a failed output."""
    values_wb = load_workbook(io.BytesIO(content), data_only=True)
    formulas_wb = load_workbook(io.BytesIO(content), data_only=False)
    ws = values_wb.active
    formula_ws = formulas_wb.active
    config = normalized_contract(template)
    columns = config["columns"]
    start = int(config["product_start"])
    expected_subtotal = Decimal(0)
    for index, product in enumerate(products, start):
        wanted = {"product_code": product["product_code"], "quantity": product["rfq_quantity"],
                  "unit": product["rfq_unit"], "unit_price": product["matched_product"].get("price")}
        for field, expected in wanted.items():
            positions = [col for col, key in columns.items() if key == field]
            if not positions:
                raise ValueError("INQUIRY_REQUIRED_COLUMN_MISSING")
            for col in positions:
                actual = ws[f"{col}{index}"].value
                if field in ("quantity", "unit_price") and expected is not None:
                    if actual is None or Decimal(str(actual)) != Decimal(str(expected)):
                        raise ValueError("INQUIRY_NUMERIC_CELL_MISMATCH")
                elif actual != expected:
                    raise ValueError("INQUIRY_SOURCE_CELL_MISMATCH")
        quantity = product.get("rfq_quantity")
        price = (product.get("matched_product") or {}).get("price")
        if quantity is not None and price is not None:
            expected_subtotal += Decimal(str(quantity)) * Decimal(str(price))
    position = (template.field_positions or {}).get("po_number")
    if position:
        position = position if isinstance(position, str) else position["position"]
        if ws[position].value != metadata.get("po_number"):
            raise ValueError("INQUIRY_PO_CELL_MISMATCH")
    if "PO Source" not in values_wb.sheetnames:
        raise ValueError("INQUIRY_SOURCE_SHEET_MISSING")
    for sheet in values_wb:
        if any(cell.data_type == "e" for row in sheet for cell in row):
            raise ValueError("INQUIRY_FORMULA_ERROR")

    if config["is_dynamic"]:
        styles = config["styles"]
        for row in range(start, start + len(products)):
            for col, formula_template in config["product_row_formulas"].items():
                expected_formula = formula_template.replace("{row}", str(row))
                if formula_ws[f"{col}{row}"].value != expected_formula:
                    raise ValueError("INQUIRY_LINE_FORMULA_MISMATCH")

        output_rows = max(len(products), 1)
        capacity = int(config["product_end"]) - start + 1
        row_delta = output_rows - capacity
        summary_refs: dict[str, str] = {}
        summary_values: dict[str, Decimal] = {}
        for item in styles.get("summary_formulas") or []:
            source = str(item.get("cell") or "")
            col = re.match(r"[A-Z]+", source).group(0)
            row = int(re.search(r"\d+", source).group(0)) + row_delta
            target = f"{col}{row}"
            if item.get("type") == "product_sum":
                expected_formula = f"=SUM({col}{start}:{col}{start + output_rows - 1})"
                summary_refs["sum_cell"] = target
                expected_value = expected_subtotal
            else:
                expected_formula = str(item.get("formula_template") or "")
                for key, ref in summary_refs.items():
                    expected_formula = expected_formula.replace(f"{{{key}}}", ref)
                label = str(item.get("label") or "").lower()
                if "tax" in label:
                    summary_refs["tax_cell"] = target
                    expected_value = expected_subtotal * Decimal("0.08")
                elif "grand" in label or ("total" in label and "sub" not in label):
                    summary_refs["grand_total_cell"] = target
                    expected_value = expected_subtotal * Decimal("1.08")
                else:
                    expected_value = Decimal(str(ws[target].value or 0))
            if formula_ws[target].value != expected_formula:
                raise ValueError("INQUIRY_SUMMARY_FORMULA_MISMATCH")
            actual_value = Decimal(str(ws[target].value or 0))
            if actual_value != expected_value:
                raise ValueError("INQUIRY_SUMMARY_VALUE_MISMATCH")
            summary_values[target] = expected_value

        for item in styles.get("external_refs") or []:
            position = str(item.get("cell") or "")
            expected_formula = str(item.get("formula_template") or "")
            for key, ref in summary_refs.items():
                expected_formula = expected_formula.replace(f"{{{key}}}", ref)
            if formula_ws[position].value != expected_formula:
                raise ValueError("INQUIRY_EXTERNAL_REFERENCE_MISMATCH")
            referenced = expected_formula.removeprefix("=").replace("$", "")
            if referenced in summary_values:
                actual_value = Decimal(str(ws[position].value or 0))
                if actual_value != summary_values[referenced]:
                    raise ValueError("INQUIRY_EXTERNAL_VALUE_MISMATCH")

        expected_summary_start = int(config["summary_start"]) + row_delta
        for position, expected in (styles.get("summary_static_values") or {}).items():
            source_row = int(re.search(r"\d+", position).group(0))
            col = re.match(r"[A-Z]+", position).group(0)
            target = f"{col}{source_row + row_delta}"
            if formula_ws[target].value != expected:
                raise ValueError("INQUIRY_SUMMARY_STRUCTURE_MISMATCH")
        if expected_summary_start != start + output_rows:
            raise ValueError("INQUIRY_ZONE_BOUNDARY_MISMATCH")
