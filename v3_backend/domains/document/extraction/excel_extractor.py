"""Excel extractor — every sheet becomes a Markdown table.

Works with .xlsx only (openpyxl). No heuristic header detection: the
first non-empty row is the header, subsequent non-empty rows are data.
"""

from __future__ import annotations

import logging
import time
from io import BytesIO

from domains.document.extraction.base import ExtractionError
from domains.document.extraction.router import register_extractor
from domains.document.extraction.schema import (
    EXTRACTION_SCHEMA_VERSION,
    ExtractedDocument,
)

logger = logging.getLogger(__name__)

# openpyxl handles modern .xlsx (OOXML) only. Legacy .xls (Microsoft CFBF
# binary) MUST be left to markitdown_extractor, which uses xlrd under the
# hood. Letting this extractor claim "application/vnd.ms-excel" caused the
# R1 2026-06-16 bug: a real Royal Caribbean .xls from a Japanese supplier
# was misrouted here and openpyxl raised "no valid workbook part" before
# markitdown ever saw it. Whitelist below intentionally excludes the legacy
# mime — markitdown_extractor._SUPPORTED_MIMES still claims it.
_SUPPORTED = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "excel",
    "xlsx",
}

# Markdown row cap per sheet — beyond this we truncate with a note. Keeps
# the rendered markdown readable + bounds DB JSON size for very long sheets.
_MAX_ROWS = 1000

# Hard scan cap — pathological ~10^6-row sheets shouldn't lock the worker.
_MAX_SCAN_ROWS = 5000


class OpenpyxlExtractor:
    name = "openpyxl"

    def supports(self, mime_type: str) -> bool:
        return mime_type in _SUPPORTED

    def extract(self, file_bytes: bytes, mime_type: str) -> ExtractedDocument:
        if not file_bytes:
            raise ExtractionError("empty Excel file", kind="empty")
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover
            raise ExtractionError(
                "openpyxl is not installed; reinstall with `pip install -e .`",
                kind="unconfigured",
            ) from exc

        start = time.perf_counter()
        try:
            wb = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
        except Exception as exc:
            raise ExtractionError(f"Failed to parse Excel: {exc}", kind="input") from exc

        try:
            sheet_count = 0
            title: str | None = None
            chunks: list[str] = []

            for ws in wb.worksheets:
                rendered = _render_sheet(ws)
                if rendered is None:
                    continue
                chunks.append(rendered)
                sheet_count += 1
                if title is None:
                    title = ws.title

            elapsed = time.perf_counter() - start
            markdown = "\n\n".join(chunks).strip()
            if not markdown:
                raise ExtractionError("no sheets with data", kind="empty")

            doc: ExtractedDocument = {
                "schema_version": EXTRACTION_SCHEMA_VERSION,
                "language": None,
                "page_count": sheet_count,
                "title": title,
                "markdown": markdown,
                "stats": {
                    "extractor": self.name,
                    "elapsed_seconds": round(elapsed, 3),
                    "input_tokens": None,
                    "output_tokens": None,
                    "finish_reason": None,
                    "truncated": False,
                },
            }
            return doc
        finally:
            wb.close()


def _render_sheet(ws) -> str | None:  # type: ignore[no-untyped-def]
    """Convert one worksheet to a Markdown table. Returns None for empty sheets.

    Header detection: the first row that has any non-blank cell is the header.
    Empty cells in subsequent rows render as "" in their column.
    """
    header_row: list[str] = []
    body_rows: list[list[str]] = []
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        row_values = ["" if v is None else str(v) for v in row]
        if not header_row and any(v.strip() for v in row_values):
            header_row = [v if v else f"col_{i + 1}" for i, v in enumerate(row_values)]
            continue
        if header_row and any(v.strip() for v in row_values):
            body_rows.append(row_values)
        if row_idx > _MAX_SCAN_ROWS:
            break

    if not header_row:
        return None

    lines = [f"## {ws.title}", ""]
    lines.append("| " + " | ".join(header_row) + " |")
    lines.append("| " + " | ".join("---" for _ in header_row) + " |")
    for row in body_rows[:_MAX_ROWS]:
        # Pad / truncate row to header width so the markdown table stays valid.
        cells = list(row[: len(header_row)]) + [""] * max(0, len(header_row) - len(row))
        lines.append("| " + " | ".join(cells) + " |")
    if len(body_rows) > _MAX_ROWS:
        lines.append(f"_(+{len(body_rows) - _MAX_ROWS} more rows)_")
    return "\n".join(lines)


register_extractor(OpenpyxlExtractor())
