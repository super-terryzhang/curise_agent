"""User-visible temporary product import journey through the real HTTP API."""

from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from domains.product_imports.template import DATA_START_ROW, PRICE_SHEET, PRODUCT_SHEET
from test_v2.integration.product_imports.test_validation import setup_data


def _login(client, seed_user) -> dict[str, str]:
    response = client.post(
        "/api/auth/login",
        json={"email": seed_user.email, "password": "password123"},
    )
    assert response.status_code == 200
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def _fill_downloaded_template(blob: bytes) -> bytes:
    workbook = load_workbook(BytesIO(blob), data_only=False)
    products = workbook[PRODUCT_SHEET]
    product_columns = {cell.value: cell.column for cell in products[4]}
    product = {
        "产品代码": "E2E-TEMP-001",
        "港口": "大阪",
        "产品名称": "TEMPORARY IMPORT APPLE",
        "供应商": "松武",
        "单位": "KG",
        "商品分类": "青果",
        "品牌": "E2E",
        "业务分类": "临时产品",
        "状态": "启用",
        "国家": "日本",
    }
    for label, value in product.items():
        products.cell(DATA_START_ROW, product_columns[label], value)

    prices = workbook[PRICE_SHEET]
    price_columns = {cell.value: cell.column for cell in prices[4]}
    periods = []
    for month, amount in ((1, 100), (2, 110), (3, 120)):
        periods.append(
            {
                "产品代码": "E2E-TEMP-001",
                "港口": "大阪",
                "价格类型": "采购价",
                "价格": amount,
                "币种": "JPY",
                "开始日期": date(2027, month, 1),
                "结束日期": date(2027, month, 28),
            }
        )
        periods.append(
            {
                "产品代码": "E2E-TEMP-001",
                "港口": "大阪",
                "价格类型": "卖价",
                "价格": amount + 50,
                "币种": "JPY",
                "开始日期": date(2027, month, 1),
                "结束日期": date(2027, month, 28),
            }
        )
    for offset, period in enumerate(periods):
        for label, value in period.items():
            prices.cell(DATA_START_ROW + offset, price_columns[label], value)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_download_upload_validate_commit_read_retry_and_rollback(
    client, db, seed_user, monkeypatch
):
    setup_data(db)
    monkeypatch.setattr(
        "apps.http.database_setup.settings.TEMP_DATABASE_SETUP_ENABLED", True
    )
    headers = _login(client, seed_user)

    initial = client.get("/api/database-setup/status", headers=headers)
    assert initial.status_code == 200
    assert initial.json()["product_count"] == 0
    assert initial.json()["price_period_count"] == 0

    template = client.get("/api/database-setup/product-template", headers=headers)
    assert template.status_code == 200
    assert "product-import-template.xlsx" in template.headers["content-disposition"]
    uploaded_workbook = _fill_downloaded_template(template.content)

    uploaded = client.post(
        "/api/database-setup/imports",
        headers=headers,
        files={
            "file": (
                "temporary-products.xlsx",
                uploaded_workbook,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert uploaded.status_code == 201, uploaded.text
    batch_id = uploaded.json()["id"]

    checked = client.post(
        f"/api/database-setup/imports/{batch_id}/validate", headers=headers
    )
    assert checked.status_code == 200, checked.text
    assert checked.json()["status"] == "ready"
    assert checked.json()["counts"] == {
        "create": 7,
        "update": 0,
        "skip": 0,
        "warning": 0,
        "block": 0,
    }

    first_page = client.get(
        f"/api/database-setup/imports/{batch_id}/rows?page=1&page_size=3",
        headers=headers,
    )
    last_page = client.get(
        f"/api/database-setup/imports/{batch_id}/rows?page=3&page_size=3",
        headers=headers,
    )
    assert first_page.status_code == last_page.status_code == 200
    assert first_page.json()["total"] == 7
    assert first_page.json()["pages"] == 3
    assert len(last_page.json()["items"]) == 1
    assert first_page.json()["can_commit"] is True

    committed = client.post(
        f"/api/database-setup/imports/{batch_id}/commit", headers=headers
    )
    retried = client.post(
        f"/api/database-setup/imports/{batch_id}/commit", headers=headers
    )
    assert committed.status_code == retried.status_code == 200
    assert committed.json() == retried.json()
    assert committed.json()["created"] == 7

    listing = client.get(
        "/api/database-setup/products?q=E2E-TEMP-001", headers=headers
    )
    assert listing.status_code == 200
    assert listing.json()["total"] == 1
    product = listing.json()["items"][0]
    assert product["port"] == "大阪"
    assert next(
        field for field in product["extensions"] if field["label"] == "业务分类"
    )["value"] == "临时产品"

    detail = client.get(
        f"/api/database-setup/products/{product['id']}", headers=headers
    )
    assert detail.status_code == 200
    periods = detail.json()["price_periods"]
    assert len(periods) == 6
    assert [item["effective_from"] for item in periods[:3]] == [
        "2027-01-01",
        "2027-02-01",
        "2027-03-01",
    ]
    assert {item["price_type"] for item in periods} == {"purchase", "selling"}

    rolled_back = client.post(
        f"/api/database-setup/imports/{batch_id}/rollback", headers=headers
    )
    rollback_retry = client.post(
        f"/api/database-setup/imports/{batch_id}/rollback", headers=headers
    )
    assert rolled_back.status_code == rollback_retry.status_code == 200
    assert rolled_back.json() == rollback_retry.json()
    assert rolled_back.json()["archived"] == 8

    after = client.get("/api/database-setup/status", headers=headers)
    assert after.status_code == 200
    assert after.json()["product_count"] == 0
    assert after.json()["price_period_count"] == 0
    empty_listing = client.get(
        "/api/database-setup/products?q=E2E-TEMP-001", headers=headers
    )
    assert empty_listing.json()["total"] == 0
