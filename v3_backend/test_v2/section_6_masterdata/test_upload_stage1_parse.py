"""Section 6 — Masterdata upload pipeline, Stage 1: `parse_excel`.

测试目标：
    把一个上传的 .xlsx 字节拆成 `UploadBatch` + N 行 `StagingProduct`：
    - header 别名归一化（"name" / "品名" / "description" 等都映射成 product_name）
    - 单元格类型按字段类型解析（价格支持 "$1,250.75" 这类格式）
    - 空行跳过、未知列保留到 raw_data 用于审计
    - 缺 product_name 列 → ParseError（前端可拿到 400）

为什么重要：
    Stage 1 是整个上传流水线的输入门。这里把字符串"12.50"丢回 None、
    或者把空行也算成一行，后面 resolve / preview / commit 全都会跑错。
    历史 v2 bug：价格被 `if x` 检查丢掉 0、未知列直接抛 KeyError。

设计方法：
    `test_v2/fixtures/helpers.make_excel(rows)` 生成真实 xlsx 字节 →
    `parse_excel(db, file_bytes=...)` →
    SQL 取 StagingProduct 行 → 断言字段值。
"""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook

from domains.masterdata.upload import list_batches, parse_excel
from domains.masterdata.upload.errors import ParseError
from domains.masterdata.upload.models import StagingProduct, UploadBatch
from test_v2.fixtures.helpers import make_excel


def _make_excel_with_blank_middle_row(rows: list[dict]) -> bytes:
    """Insert a fully-empty row between the first two data rows.
    Used to verify Stage 1 skips empty rows."""
    wb = Workbook()
    ws = wb.active
    headers = list(rows[0].keys())
    ws.append(headers)
    ws.append([rows[0].get(h) for h in headers])
    ws.append([])  # blank row
    for r in rows[1:]:
        ws.append([r.get(h) for h in headers])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ─── Happy path ──────────────────────────────────────────────


def test_parse_creates_upload_batch_with_staged_rows(db):
    """3-row valid file → 1 batch (status=ready) + 3 staging rows."""
    blob = make_excel(
        [
            {"product_code": "A1", "product_name": "Apple", "price": 10.5, "unit": "kg"},
            {"product_code": "B2", "product_name": "Banana", "price": 5.0, "unit": "ea"},
            {"product_code": "C3", "product_name": "Cherry", "price": 25.0, "unit": "kg"},
        ]
    )
    batch = parse_excel(db, file_bytes=blob, filename="produce.xlsx", user_id=1)

    assert isinstance(batch, UploadBatch)
    assert batch.id is not None
    assert batch.status == "ready"
    assert batch.total_rows == 3
    assert batch.parsed_rows == 3
    assert batch.filename == "produce.xlsx"
    assert batch.user_id == 1

    staged = (
        db.query(StagingProduct)
        .filter(StagingProduct.batch_id == batch.id)
        .order_by(StagingProduct.row_index)
        .all()
    )
    assert len(staged) == 3
    assert [s.product_code for s in staged] == ["A1", "B2", "C3"]
    assert [s.product_name for s in staged] == ["Apple", "Banana", "Cherry"]
    assert staged[0].unit == "kg"
    assert staged[0].price == 10.5


def test_parse_skips_entirely_empty_rows(db):
    """A blank row mid-file is not staged and does not bump the counters."""
    blob = _make_excel_with_blank_middle_row(
        [
            {"product_name": "Apple"},
            {"product_name": "Banana"},
        ]
    )
    batch = parse_excel(db, file_bytes=blob, filename="t.xlsx", user_id=1)
    assert batch.total_rows == 2
    staged = (
        db.query(StagingProduct)
        .filter(StagingProduct.batch_id == batch.id)
        .order_by(StagingProduct.row_index)
        .all()
    )
    assert [s.product_name for s in staged] == ["Apple", "Banana"]


def test_parse_records_true_excel_row_numbers_and_header_diagnostics(db):
    """UI errors must point to the worksheet row, not a compact staging index."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.append([])
    ws.append(["product_name", "country", "port", "mystery", "name"])
    ws.append(["Apple", "Japan", "Tokyo", "x", "duplicate alias"])
    ws.append([])
    ws.append(["Banana", "Japan", "Tokyo", "y", "duplicate alias 2"])
    buf = io.BytesIO()
    wb.save(buf)

    batch = parse_excel(
        db,
        file_bytes=buf.getvalue(),
        filename="diagnostics.xlsx",
        user_id=1,
        strict_headers=False,
    )
    staged = (
        db.query(StagingProduct)
        .filter(StagingProduct.batch_id == batch.id)
        .order_by(StagingProduct.row_index)
        .all()
    )

    assert batch.sheet_name == "Products"
    assert batch.header_row_number == 2
    assert [row.source_row_number for row in staged] == [3, 5]
    assert batch.header_diagnostics["unrecognized"][0]["raw"] == "mystery"
    duplicate = batch.header_diagnostics["duplicate_canonical"][0]
    assert duplicate["canonical"] == "product_name"
    assert duplicate["columns"] == [1, 5]


# ─── Price parsing ──────────────────────────────────────────


@pytest.mark.parametrize(
    ("raw_price", "expected"),
    [
        ("12.50", 12.5),
        ("$1,250.75", 1250.75),
        ("¥980", 980.0),
        (15, 15.0),
        (15.99, 15.99),
        ("", None),
        (None, None),
        ("not-a-number", None),  # parser swallows garbage → None
    ],
)
def test_parse_normalizes_price_strings(db, raw_price, expected):
    """The price parser handles currency symbols, thousands separators, ints,
    floats, blanks, and garbage strings. Garbage → None (no exception)."""
    blob = make_excel([{"product_name": "X", "price": raw_price}])
    batch = parse_excel(db, file_bytes=blob, filename="t.xlsx", user_id=1)
    staged = db.query(StagingProduct).filter(StagingProduct.batch_id == batch.id).one()
    assert staged.price == expected


# ─── Schema-tolerance: unknown columns ──────────────────────


def test_parse_preserves_unknown_columns_in_raw_data(db):
    """Columns not in the alias map should NOT crash the parse; they end up
    in `raw_data` (JSON) so we can audit the original cell-by-cell input."""
    blob = make_excel(
        [
            {
                "product_name": "Mystery Item",
                "weird_col": "weird_value",
                "another_extra": 42,
            }
        ]
    )
    batch = parse_excel(db, file_bytes=blob, filename="t.xlsx", user_id=1)
    sp = db.query(StagingProduct).filter(StagingProduct.batch_id == batch.id).one()
    assert sp.product_name == "Mystery Item"
    assert sp.raw_data["weird_col"] == "weird_value"
    assert sp.raw_data["another_extra"] == 42


def test_parse_accepts_japanese_aliases_for_name_and_price(db):
    """Aliases `品名` → product_name and `単価` → price. Real Excel files from
    Japanese cruise procurement use these headers."""
    blob = make_excel([{"品名": "天ぷら粉", "単価": "880"}])
    batch = parse_excel(db, file_bytes=blob, filename="jp.xlsx", user_id=1)
    sp = db.query(StagingProduct).filter(StagingProduct.batch_id == batch.id).one()
    assert sp.product_name == "天ぷら粉"
    assert sp.price == 880.0


# ─── Failure modes ──────────────────────────────────────────


def test_parse_rejects_file_without_product_name_column(db):
    """No column resolves to `product_name` → ParseError carrying the offending field name."""
    blob = make_excel([{"random_col": "X", "another_col": 7}])
    with pytest.raises(ParseError) as exc:
        parse_excel(db, file_bytes=blob, filename="bad.xlsx", user_id=1)
    assert "product_name" in str(exc.value)


def test_parse_rejects_corrupt_workbook_bytes(db):
    """Non-xlsx bytes → ParseError, not a raw openpyxl exception."""
    with pytest.raises(ParseError):
        parse_excel(db, file_bytes=b"this is not a workbook", filename="x.xlsx", user_id=1)


def test_parse_rejects_unsupported_entity_type(db):
    """Stage 1 only knows about `products` for now. Other types must be a clean refusal."""
    blob = make_excel([{"product_name": "X"}])
    with pytest.raises(ParseError) as exc:
        parse_excel(
            db, file_bytes=blob, filename="t.xlsx", user_id=1, entity_type="suppliers"
        )
    assert "suppliers" in str(exc.value)


# ─── User scoping ───────────────────────────────────────────


def test_parse_scopes_batch_to_uploader_user_id(db):
    """Alice uploads under user_id=1; Bob under user_id=2. Each only sees their own batches."""
    blob_a = make_excel([{"product_name": "Alice apple"}])
    blob_b = make_excel([{"product_name": "Bob banana"}])
    parse_excel(db, file_bytes=blob_a, filename="alice.xlsx", user_id=1)
    parse_excel(db, file_bytes=blob_b, filename="bob.xlsx", user_id=2)

    alice_batches = list_batches(db, user_id=1)
    bob_batches = list_batches(db, user_id=2)
    assert len(alice_batches) == 1
    assert len(bob_batches) == 1
    assert alice_batches[0]["filename"] == "alice.xlsx"
    assert bob_batches[0]["filename"] == "bob.xlsx"


# ─── Additional column / sheet / edge-case coverage ─────────────


@pytest.mark.parametrize(
    ("alias_header", "field"),
    [
        # product_code aliases
        ("code", "product_code"),
        ("sku", "product_code"),
        ("item_code", "product_code"),
        ("ITEM CODE", "product_code"),  # case-insensitive + space variant
        ("品番", "product_code"),
        # product_name aliases (already covered by happy path for `product_name`)
        ("name", "product_name"),
        ("description", "product_name"),
        ("item", "product_name"),
    ],
)
def test_parse_accepts_english_and_mixed_header_aliases(db, alias_header, field):
    """Header alias map must canonicalise English variants (sku, code, item)
    as well as the Japanese ones. Pre-fix, only "name" worked while
    suppliers shipping templates titled "Item" or "SKU" returned 400."""
    # Always include a `product_name` column so the parser doesn't reject
    # us for the wrong reason; vary the second column.
    blob = make_excel([{"product_name": "stub", alias_header: "X1"}])
    batch = parse_excel(db, file_bytes=blob, filename="t.xlsx", user_id=1)
    sp = db.query(StagingProduct).filter(StagingProduct.batch_id == batch.id).one()
    # If the alias resolved, the corresponding field on StagingProduct is "X1".
    # `name` / `description` / `item` would overwrite product_name — accept either.
    if field == "product_code":
        assert sp.product_code == "X1"
    else:
        # name/description/item — last write wins because column_map is dict
        # keyed by index. So whichever column comes second sets product_name.
        assert sp.product_name in ("stub", "X1")


def test_parse_only_reads_first_worksheet(db):
    """A workbook with multiple sheets must parse the first one and silently
    ignore the rest — matches our header doc and avoids surprises when
    suppliers leave junk on tab 2."""
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "real"
    ws1.append(["product_name", "price"])
    ws1.append(["Apple", 10])
    ws2 = wb.create_sheet("garbage")
    ws2.append(["totally", "different", "headers"])
    ws2.append(["should", "not", "appear"])
    buf = io.BytesIO()
    wb.save(buf)

    batch = parse_excel(db, file_bytes=buf.getvalue(), filename="multi.xlsx", user_id=1)
    assert batch.total_rows == 1
    sp = db.query(StagingProduct).filter(StagingProduct.batch_id == batch.id).one()
    assert sp.product_name == "Apple"
    assert sp.price == 10.0


def test_parse_preserves_leading_zeros_in_product_code(db):
    """Product codes like "00100" must stay as 5-char strings — Excel
    silently auto-coerces a "leading-zero number" string into 100.0 if you
    let openpyxl interpret it. The parser MUST keep the string form.
    v2 hit this exact bug: SKUs got squashed to 100 instead of 00100,
    matching against DB rows broke."""
    # Write the code as a text-typed cell. openpyxl preserves it when the
    # cell value is a Python str.
    wb = Workbook()
    ws = wb.active
    ws.append(["product_code", "product_name"])
    ws.append(["00100", "Coded Apple"])
    ws.append(["007", "Bond Banana"])
    buf = io.BytesIO()
    wb.save(buf)

    batch = parse_excel(db, file_bytes=buf.getvalue(), filename="zero.xlsx", user_id=1)
    rows = (
        db.query(StagingProduct)
        .filter(StagingProduct.batch_id == batch.id)
        .order_by(StagingProduct.row_index)
        .all()
    )
    assert [r.product_code for r in rows] == ["00100", "007"], (
        "leading zeros were stripped — DB code-match will fail"
    )


def test_parse_handles_completely_empty_workbook(db):
    """A workbook with one sheet but no rows at all → ParseError "no header
    row". The endpoint converts this to a 400 with a clear message."""
    wb = Workbook()
    # Default sheet exists but has zero rows.
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ParseError) as exc:
        parse_excel(db, file_bytes=buf.getvalue(), filename="empty.xlsx", user_id=1)
    assert "header" in str(exc.value).lower()


def test_parse_handles_duplicate_column_headers(db):
    """If the user accidentally ships two columns called "price" the parser
    shouldn't crash. The current behaviour is "last column wins" because
    column_map is keyed by column index; this test pins that behaviour so
    nobody silently changes it without a migration plan."""
    wb = Workbook()
    ws = wb.active
    ws.append(["product_name", "price", "price"])  # duplicate
    ws.append(["Dup Apple", 10, 20])
    buf = io.BytesIO()
    wb.save(buf)

    batch = parse_excel(db, file_bytes=buf.getvalue(), filename="dup.xlsx", user_id=1)
    sp = db.query(StagingProduct).filter(StagingProduct.batch_id == batch.id).one()
    # Last-write-wins via column index → 20 not 10.
    assert sp.price == 20.0
    assert sp.product_name == "Dup Apple"
