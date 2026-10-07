"""Cross-sheet and database-aware product import validation."""

from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from domains.dynamic_data.models import DataTable
from domains.dynamic_data.schemas import Actor
from domains.masterdata.models import Category, Country, Port, Product, ProductPricePeriod, Supplier
from domains.product_imports.parser import parse_workbook
from domains.product_imports.template import (
    DATA_START_ROW,
    PRICE_SHEET,
    PRODUCT_SHEET,
    build_product_workbook,
)
from domains.product_imports.validation import validate_batch
from scripts.seed_clean_product_fields import PRODUCT_TABLE_ID, seed_product_business_classification

ADMIN = Actor(id=1, role="admin")


def setup_data(db):
    table = DataTable(
        id=PRODUCT_TABLE_ID,
        name="产品",
        table_kind="system",
        system_key="products",
        status="active",
        schema_version=1,
        created_by=0,
        updated_by=0,
    )
    db.add(table)
    db.commit()
    seed_product_business_classification(db, actor_id=0)
    japan = Country(name="日本", code="JPN", status=True)
    db.add(japan)
    db.flush()
    osaka = Port(name="大阪", country_id=japan.id, status=True)
    tokyo = Port(name="东京", country_id=japan.id, status=True)
    supplier = Supplier(name="松武", status=True)
    category = Category(name="青果", status=True)
    db.add_all([osaka, tokyo, supplier, category])
    db.commit()
    return japan, osaka, tokyo, supplier, category


def build_upload(db, product_rows=(), price_rows=()):
    blob = build_product_workbook(db, ADMIN, include_existing=False)
    wb = load_workbook(BytesIO(blob), data_only=False)
    for sheet_name, rows in ((PRODUCT_SHEET, product_rows), (PRICE_SHEET, price_rows)):
        sheet = wb[sheet_name]
        columns = {cell.value: cell.column for cell in sheet[4]}
        for offset, values in enumerate(rows):
            for label, value in values.items():
                sheet.cell(DATA_START_ROW + offset, columns[label], value)
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def validate_upload(db, product_rows=(), price_rows=()):
    batch = parse_workbook(
        db,
        build_upload(db, product_rows, price_rows),
        "产品导入.xlsx",
        user_id=1,
    )
    return validate_batch(db, batch.id, user_id=1)


def product(code, port="大阪", **values):
    return {
        "产品代码": code,
        "港口": port,
        "产品名称": values.pop("name", "APPLE"),
        "国家": values.pop("country", "日本"),
        "状态": values.pop("status", "启用"),
        **values,
    }


def price(code, start, end, *, port="大阪", kind="采购价", amount=100, currency="JPY"):
    return {
        "产品代码": code,
        "港口": port,
        "价格类型": kind,
        "价格": amount,
        "币种": currency,
        "开始日期": start,
        "结束日期": end,
    }


def codes(preview):
    return {issue.code for row in preview.rows for issue in row.issues}


def test_new_product_can_reference_same_batch_prices_and_missing_type_is_warning(db):
    setup_data(db)
    preview = validate_upload(
        db,
        [product(" P-001 ", 业务分类="中标产品")],
        [price("p-001", date(2026, 1, 1), date(2026, 1, 31))],
    )

    assert preview.status == "ready"
    assert preview.counts.create == 2
    assert preview.counts.block == 0
    assert "MISSING_SELLING_PRICE" in codes(preview)
    product_row = next(row for row in preview.rows if row.sheet == PRODUCT_SHEET)
    assert product_row.product_key == "p-001@1"
    assert product_row.normalized_values["extension_values"]


def test_same_code_different_port_is_distinct_but_conflicting_duplicate_rows_block(db):
    setup_data(db)
    distinct = validate_upload(
        db,
        [product("ABC", "大阪"), product(" abc ", "东京", name="TOKYO APPLE")],
    )
    assert distinct.counts.create == 2
    assert distinct.counts.block == 0

    conflict = validate_upload(
        db,
        [product("ABC", "大阪", name="ONE"), product(" abc ", "大阪", name="TWO")],
    )
    assert conflict.status == "failed"
    assert conflict.counts.block == 2
    duplicate_issues = [
        issue for row in conflict.rows for issue in row.issues if issue.code == "DUPLICATE_PRODUCT"
    ]
    assert {tuple(issue.related_rows) for issue in duplicate_issues} == {
        ((PRODUCT_SHEET, DATA_START_ROW), (PRODUCT_SHEET, DATA_START_ROW + 1))
    }


def test_unknown_master_and_inactive_choice_are_specific_blocking_issues(db):
    setup_data(db)
    preview = validate_upload(
        db,
        [product("P-001", "不存在", 供应商="无此供应商", 业务分类="不存在选项")],
    )
    assert preview.status == "failed"
    assert {"UNKNOWN_PORT", "UNKNOWN_SUPPLIER", "UNKNOWN_OPTION"}.issubset(codes(preview))


def test_existing_sparse_update_clear_marker_and_exact_period_decisions(db):
    japan, osaka, _, supplier, category = setup_data(db)
    existing = Product(
        product_name_en="APPLE",
        code="P-001",
        port_id=osaka.id,
        country_id=japan.id,
        supplier_id=supplier.id,
        category_id=category.id,
        unit="KG",
        brand="OLD",
        status=True,
    )
    db.add(existing)
    db.flush()
    period = ProductPricePeriod(
        product_id=existing.id,
        price_type="purchase",
        amount=Decimal("100"),
        currency="JPY",
        effective_from=date(2026, 1, 1),
        effective_to=date(2026, 1, 31),
    )
    db.add(period)
    db.commit()

    preview = validate_upload(
        db,
        [product("p-001", 单位="", 品牌="__CLEAR__")],
        [
            price("P-001", date(2026, 1, 1), date(2026, 1, 31), amount=110),
            price(
                "P-001",
                date(2026, 2, 1),
                date(2026, 2, 28),
                kind="卖价",
                amount=180,
            ),
        ],
    )
    product_row = next(row for row in preview.rows if row.sheet == PRODUCT_SHEET)
    purchase = next(
        row
        for row in preview.rows
        if row.sheet == PRICE_SHEET and row.normalized_values.get("price_type") == "purchase"
    )
    assert product_row.action == "update"
    assert product_row.normalized_values["core_values"]["unit"] == "KG"
    assert product_row.normalized_values["core_values"]["brand"] is None
    assert purchase.action == "update"
    assert purchase.target_id == period.id


def test_file_and_database_partial_overlaps_block_and_missing_product_blocks(db):
    japan, osaka, _, _, _ = setup_data(db)
    existing = Product(
        product_name_en="APPLE",
        code="P-001",
        port_id=osaka.id,
        country_id=japan.id,
        status=True,
    )
    db.add(existing)
    db.flush()
    db.add(
        ProductPricePeriod(
            product_id=existing.id,
            price_type="purchase",
            amount=Decimal("100"),
            currency="JPY",
            effective_from=date(2026, 1, 1),
            effective_to=date(2026, 1, 31),
        )
    )
    db.commit()

    preview = validate_upload(
        db,
        [],
        [
            price("P-001", date(2026, 1, 20), date(2026, 2, 10)),
            price("P-001", date(2026, 2, 5), date(2026, 2, 28)),
            price("MISSING", date(2026, 3, 1), date(2026, 3, 31)),
        ],
    )
    assert preview.status == "failed"
    assert {"DATABASE_PERIOD_OVERLAP", "FILE_PERIOD_OVERLAP", "PRODUCT_NOT_FOUND"}.issubset(
        codes(preview)
    )


def test_exported_ids_make_unchanged_rows_skip_and_stale_versions_block(db):
    japan, osaka, _, _, _ = setup_data(db)
    existing = Product(
        product_name_en="APPLE",
        code="P-001",
        port_id=osaka.id,
        country_id=japan.id,
        brand="OLD",
        status=True,
    )
    db.add(existing)
    db.flush()
    period = ProductPricePeriod(
        product_id=existing.id,
        price_type="purchase",
        amount=Decimal("100"),
        currency="JPY",
        effective_from=date(2026, 1, 1),
        effective_to=date(2026, 1, 31),
    )
    db.add(period)
    db.commit()
    exported = build_product_workbook(
        db, ADMIN, include_existing=True, product_ids=[existing.id]
    )

    unchanged_batch = parse_workbook(db, exported, "原样.xlsx", user_id=1)
    unchanged = validate_batch(db, unchanged_batch.id, user_id=1)
    assert [row.action for row in unchanged.rows] == ["skip", "skip"]

    existing.brand = "CHANGED ELSEWHERE"
    period.amount = Decimal("120")
    db.commit()
    stale_batch = parse_workbook(db, exported, "过期.xlsx", user_id=1)
    stale = validate_batch(db, stale_batch.id, user_id=1)
    assert stale.status == "failed"
    assert {"PRODUCT_CHANGED", "PERIOD_CHANGED"}.issubset(codes(stale))
