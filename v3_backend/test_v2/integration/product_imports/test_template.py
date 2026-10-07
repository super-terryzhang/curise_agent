"""Schema-versioned workbook generation for product data and price periods."""

import json
from datetime import date
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook

from domains.dynamic_data import service as data_service
from domains.dynamic_data.models import DataField, DataRecord, DataTable
from domains.dynamic_data.schemas import Actor, FieldCreate
from domains.masterdata.models import (
    Category,
    Country,
    ExchangeRate,
    Port,
    Product,
    ProductPricePeriod,
    Supplier,
)
from domains.product_imports.template import (
    DATA_START_ROW,
    PRICE_SHEET,
    PRODUCT_SHEET,
    SYSTEM_SHEET,
    build_product_workbook,
)
from scripts.seed_clean_product_fields import (
    BUSINESS_CLASSIFICATION_FIELD_ID,
    BUSINESS_CLASSIFICATION_OPTIONS,
    PRODUCT_TABLE_ID,
    seed_product_business_classification,
)

ADMIN = Actor(id=1, role="admin")


def install_product_catalog(db):
    table = DataTable(
        id=PRODUCT_TABLE_ID,
        name="产品",
        table_kind="system",
        system_key="products",
        status="active",
        schema_version=1,
        created_by=0,
        updated_by=0,
    )
    db.add(table)
    db.commit()
    seed_product_business_classification(db, actor_id=0)
    db.commit()
    return table


def workbook(blob):
    return load_workbook(BytesIO(blob), data_only=False)


def headers(sheet):
    return [cell.value for cell in sheet[4]]


def validation_for(sheet, column_letter):
    return next(
        validation
        for validation in sheet.data_validations.dataValidation
        if f"{column_letter}{DATA_START_ROW}" in str(validation.sqref)
    )


def test_blank_template_has_two_business_sheets_hidden_manifest_and_formats(db):
    table = install_product_catalog(db)

    wb = workbook(build_product_workbook(db, ADMIN, include_existing=False))

    assert wb.sheetnames == [PRODUCT_SHEET, PRICE_SHEET, SYSTEM_SHEET]
    assert wb[SYSTEM_SHEET].sheet_state == "veryHidden"
    manifest = json.loads(wb[SYSTEM_SHEET]["A2"].value)
    assert manifest["contract_version"] == 1
    assert manifest["product_schema_version"] == table.schema_version == 2
    classification = next(
        field
        for field in manifest["product_fields"]
        if field["key"] == f"extension:{BUSINESS_CLASSIFICATION_FIELD_ID}"
    )
    assert classification == {
        "key": f"extension:{BUSINESS_CLASSIFICATION_FIELD_ID}",
        "label": "业务分类",
        "field_type": "single_select",
        "required": False,
        "option_ids": [str(option_id) for option_id, _ in BUSINESS_CLASSIFICATION_OPTIONS],
    }
    assert headers(wb[PRODUCT_SHEET]) == [
        "产品代码",
        "港口",
        "产品名称",
        "供应商",
        "单位",
        "商品分类",
        "品牌",
        "业务分类",
        "状态",
        "国家",
        "__产品ID",
        "__产品版本",
        "__扩展版本",
    ]
    assert headers(wb[PRICE_SHEET]) == [
        "产品代码",
        "港口",
        "价格类型",
        "价格",
        "币种",
        "开始日期",
        "结束日期",
        "__价格区间ID",
        "__价格区间版本",
    ]
    assert "每个产品一行" in wb[PRODUCT_SHEET]["A2"].value
    assert "每个价格区间一行" in wb[PRICE_SHEET]["A2"].value
    assert wb[PRODUCT_SHEET][f"A{DATA_START_ROW}"].number_format == "@"
    assert wb[PRICE_SHEET][f"D{DATA_START_ROW}"].number_format == "0.00"
    assert wb[PRICE_SHEET][f"F{DATA_START_ROW}"].number_format == "yyyy-mm-dd"
    assert wb[PRODUCT_SHEET].column_dimensions["K"].hidden is True
    assert wb[PRICE_SHEET].column_dimensions["H"].hidden is True
    assert wb[PRODUCT_SHEET].protection.sheet is True
    assert wb[PRODUCT_SHEET][f"A{DATA_START_ROW}"].protection.locked is False


def test_template_uses_configured_choices_and_does_not_invent_currencies(db):
    install_product_catalog(db)
    country = Country(name="日本", code="JPN", status=True)
    db.add(country)
    db.flush()
    db.add_all(
        [
            Port(name="大阪", country_id=country.id, status=True),
            Supplier(name="松武", status=True),
            Category(name="青果", status=True),
        ]
    )
    db.commit()

    wb = workbook(build_product_workbook(db, ADMIN, include_existing=False))
    product = wb[PRODUCT_SHEET]
    price = wb[PRICE_SHEET]
    assert validation_for(product, "B").type == "list"
    assert validation_for(product, "H").type == "list"
    assert validation_for(product, "I").type == "list"
    assert validation_for(price, "C").type == "list"
    currency = validation_for(price, "E")
    assert currency.type == "custom"
    assert "UPPER" in currency.formula1

    db.add(
        ExchangeRate(
            from_currency="JPY",
            to_currency="USD",
            rate=Decimal("0.01"),
            effective_date=date(2026, 10, 7),
        )
    )
    db.commit()
    wb = workbook(build_product_workbook(db, ADMIN, include_existing=False))
    currency = validation_for(wb[PRICE_SHEET], "E")
    assert currency.type == "list"
    assert {wb[SYSTEM_SHEET]["J5"].value, wb[SYSTEM_SHEET]["J6"].value} == {"JPY", "USD"}


def test_existing_export_includes_ids_extension_values_and_chronological_periods(db):
    install_product_catalog(db)
    country = Country(name="日本", code="JPN", status=True)
    supplier = Supplier(name="松武", status=True)
    category = Category(name="青果", status=True)
    db.add_all([country, supplier, category])
    db.flush()
    port = Port(name="大阪", country_id=country.id, status=True)
    db.add(port)
    db.flush()
    product = Product(
        product_name_en="APPLE",
        code="P-001",
        port_id=port.id,
        country_id=country.id,
        supplier_id=supplier.id,
        category_id=category.id,
        unit="KG",
        brand="TEST",
        status=True,
    )
    db.add(product)
    db.flush()
    selected = str(BUSINESS_CLASSIFICATION_OPTIONS[0][0])
    db.add(
        DataRecord(
            table_id=PRODUCT_TABLE_ID,
            source_record_id=str(product.id),
            values={str(BUSINESS_CLASSIFICATION_FIELD_ID): selected},
            revision=3,
            schema_version=2,
            status="active",
            created_by=1,
            updated_by=1,
        )
    )
    purchase_period = ProductPricePeriod(
                product_id=product.id,
                price_type="purchase",
                amount=Decimal("120"),
                currency="JPY",
                effective_from=date(2026, 4, 1),
                effective_to=date(2026, 4, 30),
                revision=2,
            )
    selling_period = ProductPricePeriod(
                product_id=product.id,
                price_type="selling",
                amount=Decimal("180"),
                currency="JPY",
                effective_from=date(2026, 1, 1),
                effective_to=date(2026, 3, 31),
                revision=4,
            )
    db.add_all([purchase_period, selling_period])
    db.commit()

    wb = workbook(
        build_product_workbook(db, ADMIN, include_existing=True, product_ids=[product.id])
    )
    product_values = [cell.value for cell in wb[PRODUCT_SHEET][DATA_START_ROW]]
    assert product_values[:10] == [
        "P-001",
        "大阪",
        "APPLE",
        "松武",
        "KG",
        "青果",
        "TEST",
        "中标产品",
        "启用",
        "日本",
    ]
    assert product_values[10:] == [product.id, product.revision, 3]
    price_rows = [
        [cell.value for cell in row]
        for row in wb[PRICE_SHEET].iter_rows(min_row=DATA_START_ROW, max_row=DATA_START_ROW + 1)
    ]
    assert [row[5].date() for row in price_rows] == [date(2026, 1, 1), date(2026, 4, 1)]
    assert price_rows[0][:5] == [
        "P-001",
        "大阪",
        "卖价",
        180,
        "JPY",
    ]
    assert price_rows[0][5].date() == date(2026, 1, 1)
    assert price_rows[0][6].date() == date(2026, 3, 31)
    assert price_rows[0][7] is not None
    assert price_rows[0][8] == selling_period.revision


def test_next_template_adds_active_extension_and_removes_archived_extension(db):
    table = install_product_catalog(db)
    field = data_service.create_field(
        db,
        table.id,
        FieldCreate(
            id=uuid4(),
            label="内部等级",
            field_type="text",
            expected_schema_version=db.get(DataTable, table.id).schema_version,
        ),
        actor=ADMIN,
    )
    assert "内部等级" in headers(
        workbook(build_product_workbook(db, ADMIN, include_existing=False))[PRODUCT_SHEET]
    )

    db.get(DataField, field.id).status = "archived"
    db.commit()
    assert "内部等级" not in headers(
        workbook(build_product_workbook(db, ADMIN, include_existing=False))[PRODUCT_SHEET]
    )
