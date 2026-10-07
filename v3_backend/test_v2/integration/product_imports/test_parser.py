"""Workbook parsing is bounded, structural and source-coordinate preserving."""

import json
from io import BytesIO

from openpyxl import load_workbook

from domains.dynamic_data.models import DataTable
from domains.dynamic_data.schemas import Actor
from domains.product_imports.models import ImportRow
from domains.product_imports.parser import MAX_WORKBOOK_BYTES, parse_workbook
from domains.product_imports.template import (
    DATA_START_ROW,
    PRICE_SHEET,
    PRODUCT_SHEET,
    build_product_workbook,
)
from scripts.seed_clean_product_fields import PRODUCT_TABLE_ID, seed_product_business_classification

ADMIN = Actor(id=1, role="admin")


def install_catalog(db):
    table = DataTable(
        id=PRODUCT_TABLE_ID,
        name="产品",
        table_kind="system",
        system_key="products",
        status="active",
        schema_version=1,
        created_by=0,
        updated_by=0,
    )
    db.add(table)
    db.commit()
    seed_product_business_classification(db, actor_id=0)
    db.commit()
    return table


def edited(blob, edit):
    wb = load_workbook(BytesIO(blob), data_only=False)
    edit(wb)
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def issue_codes(batch):
    return {issue["code"] for issue in (batch.result or {}).get("issues", [])}


def test_valid_workbook_stages_nonempty_rows_and_allows_empty_price_sheet(db):
    install_catalog(db)
    blob = build_product_workbook(db, ADMIN, include_existing=False)

    def fill(wb):
        sheet = wb[PRODUCT_SHEET]
        sheet.cell(DATA_START_ROW, 1, " P-001 ")
        sheet.cell(DATA_START_ROW, 2, "大阪")
        sheet.cell(DATA_START_ROW, 3, "APPLE")

    batch = parse_workbook(db, edited(blob, fill), "产品准备.xlsx", user_id=1)
    rows = db.query(ImportRow).filter(ImportRow.batch_id == batch.id).all()

    assert batch.status == "uploaded"
    assert batch.total_rows == 1
    assert [(row.sheet_key, row.source_row_number, row.action) for row in rows] == [
        ("products", DATA_START_ROW, "pending")
    ]
    assert rows[0].raw_values["product_code"] == " P-001 "


def test_parser_reports_sheet_header_formula_and_stale_schema_coordinates(db):
    table = install_catalog(db)
    original = build_product_workbook(db, ADMIN, include_existing=False)

    missing = parse_workbook(
        db,
        edited(original, lambda wb: wb[PRICE_SHEET].__setattr__("title", "价格")),
        "缺少Sheet.xlsx",
        user_id=1,
    )
    assert missing.status == "failed"
    assert "MISSING_SHEET" in issue_codes(missing)

    duplicate = parse_workbook(
        db,
        edited(
            original,
            lambda wb: wb[PRODUCT_SHEET].cell(4, 2, wb[PRODUCT_SHEET].cell(4, 1).value),
        ),
        "重复表头.xlsx",
        user_id=1,
    )
    assert {"DUPLICATE_HEADER", "MISSING_HEADER"}.issubset(issue_codes(duplicate))

    def formula_edit(wb):
        sheet = wb[PRODUCT_SHEET]
        sheet.cell(DATA_START_ROW, 1, "P-001")
        sheet.cell(DATA_START_ROW, 2, "大阪")
        sheet.cell(DATA_START_ROW, 3, "=UPPER(\"apple\")")

    formula = parse_workbook(db, edited(original, formula_edit), "公式.xlsx", user_id=1)
    staged = db.query(ImportRow).filter(ImportRow.batch_id == formula.id).one()
    assert staged.action == "block"
    assert staged.issues == [
        {
            "severity": "block",
            "code": "FORMULA_NOT_ALLOWED",
            "message": "不允许使用公式，请粘贴计算后的值",
            "sheet": PRODUCT_SHEET,
            "row": DATA_START_ROW,
            "field": "产品名称",
        }
    ]

    table.schema_version += 1
    db.commit()
    stale = parse_workbook(db, original, "旧模板.xlsx", user_id=1)
    assert stale.status == "uploaded"
    assert "STALE_SCHEMA" in issue_codes(stale)
    assert stale.block_count == 1

    def stale_contract(wb):
        payload = json.loads(wb["_系统信息"]["A2"].value)
        payload["contract_version"] = 2
        wb["_系统信息"]["A2"] = json.dumps(payload)

    unsupported = parse_workbook(
        db, edited(original, stale_contract), "旧契约.xlsx", user_id=1
    )
    assert unsupported.status == "failed"
    assert "STALE_CONTRACT" in issue_codes(unsupported)


def test_malformed_and_oversized_workbooks_become_safe_chinese_batch_issues(db):
    install_catalog(db)

    malformed = parse_workbook(db, b"not-an-xlsx", "损坏.xlsx", user_id=1)
    assert malformed.status == "failed"
    assert malformed.error_message == "无法读取 Excel 工作簿，请重新下载模板后填写"
    assert "MALFORMED_WORKBOOK" in issue_codes(malformed)

    oversized = parse_workbook(db, b"x" * (MAX_WORKBOOK_BYTES + 1), "过大.xlsx", user_id=1)
    assert oversized.status == "failed"
    assert "FILE_TOO_LARGE" in issue_codes(oversized)
