"""Saved workbook verification, including viewers without formula recalculation."""

import io
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from domains.inquiry.workbook_quality import formula_values, save_with_formula_cache


def test_saved_formula_and_cached_total_agree():
    wb = Workbook()
    ws = wb.active
    ws["H19"], ws["K19"] = 1, 7980
    ws["M19"] = "=H19*K19"
    ws["M20"] = "=H20*K20"
    ws["M102"] = "=SUM(M19:M101)"
    ws["M103"] = "=M102*0.08"
    ws["M104"] = "=M102+M103"
    ws["I13"] = "=M104"
    content = save_with_formula_cache(wb)
    values = load_workbook(io.BytesIO(content), data_only=True).active
    assert Decimal(str(values["I13"].value)) == Decimal("8618.40")
    assert values["M19"].value == 7980
    formulas = load_workbook(io.BytesIO(content), data_only=False).active
    assert formulas["M19"].value == "=H19*K19"


@pytest.mark.parametrize(
    "formula", ['=WEBSERVICE("https://invalid.example")', "=A1", "=1/0", "=SUM(A1:XFD999999)"]
)
def test_unknown_or_unsafe_formula_fails(formula):
    ws = Workbook().active
    ws["A1"] = formula
    with pytest.raises((ValueError, ArithmeticError, SyntaxError)):
        formula_values(ws)


def test_nonnumeric_arithmetic_is_not_reported_as_zero():
    ws = Workbook().active
    ws["A1"], ws["A2"], ws["A3"] = "確認中", 10, "=A1*A2"
    with pytest.raises(ValueError, match="NON_NUMERIC"):
        formula_values(ws)
