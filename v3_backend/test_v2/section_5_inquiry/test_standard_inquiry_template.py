"""Regression coverage for the contract-backed Japanese inquiry workbook."""

from __future__ import annotations

import io
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from domains.inquiry.standard_template import (
    STANDARD_TEMPLATE_ASSET,
    standard_template_definition,
)
from domains.inquiry.template_contract import (
    assert_source_workbook_matches,
    configuration_errors,
)
from domains.inquiry.template_engine import render_inquiry_excel
from domains.inquiry.workbook_quality import validate_output

TEMPLATE_PATH = STANDARD_TEMPLATE_ASSET


def _template():
    return SimpleNamespace(**standard_template_definition())


def _metadata() -> dict:
    return {
        "ship_name": "MITSUI OCEAN FUJI",
        "ship_port_title": "MITSUI OCEAN FUJI［横浜］",
        "supplier_name": "日本海洋物産株式会社",
        "supplier_contact": "田中",
        "supplier_address": "東京都中央区",
        "supplier_tel": "03-1234-5678",
        "supplier_email": "sales@example.jp",
        "generated_date": "2026-09-16",
        "inquiry_reference": "RFQ-000001-0001",
        "voyage": "V-123",
        "delivery_date": "2026-09-20",
        "delivery_time_notes": "10:00",
        "delivery_address": "横浜港 大さん橋",
        "destination_port": "横浜",
        "delivery_contact": "佐藤",
        "payment_date": "月末締翌月末",
        "payment_method": "銀行振込",
        "currency": "JPY",
        "po_number": "PO-001",
    }


def _products(count: int) -> list[dict]:
    return [
        {
            "product_code": f"SKU-{number:03}",
            "product_name": f"Item {number}",
            "product_name_en": f"Item {number}",
            "product_name_jp": f"商品 {number}",
            "description": "10 KG",
            "rfq_quantity": number,
            "rfq_unit": "CTN",
            "source_po_number": f"PO-{1 + (number % 2):03}",
            "matched_product": {
                "price": 100 + number,
                "currency": "JPY",
                "product_name_en": f"Item {number}",
                "product_name_jp": f"商品 {number}",
                "unit": "CTN",
                "pack_size": "10 KG",
            },
        }
        for number in range(1, count + 1)
    ]


@pytest.mark.parametrize("count", [1, 6, 7, 11, 12, 25])
def test_dynamic_rows_and_totals_cover_every_product(count: int) -> None:
    template = _template()
    products = _products(count)
    content = render_inquiry_excel(
        template,
        metadata=_metadata(),
        products=products,
        supplier_id=1,
        template_file_bytes=TEMPLATE_PATH.read_bytes(),
    )

    validate_output(content, template, products, _metadata())
    formulas = load_workbook(io.BytesIO(content), data_only=False).active
    values = load_workbook(io.BytesIO(content), data_only=True).active
    summary_row = 22 + count
    subtotal = sum(item["rfq_quantity"] * item["matched_product"]["price"] for item in products)

    assert formulas[f"L{summary_row}"].value == f"=SUM(L22:L{summary_row - 1})"
    assert values[f"L{summary_row}"].value == subtotal
    assert values["H16"].value == pytest.approx(subtotal * 1.08)
    assert formulas[f"L{summary_row - 1}"].value == (
        f"=H{summary_row - 1}*J{summary_row - 1}"
    )
    assert f"F{summary_row - 1}:G{summary_row - 1}" in {
        str(item) for item in formulas.merged_cells.ranges
    }


def test_headers_are_replaced_and_formatted_without_sample_leakage() -> None:
    template = _template()
    content = render_inquiry_excel(
        template,
        metadata=_metadata(),
        products=_products(2),
        supplier_id=1,
        template_file_bytes=TEMPLATE_PATH.read_bytes(),
    )
    sheet = load_workbook(io.BytesIO(content), data_only=False).active

    assert sheet["A1"].value == "PURCHASE ORDER"
    assert sheet["A2"].value == "MITSUI OCEAN FUJI［横浜］"
    assert sheet["J5"].value == "RFQ No.："
    assert sheet["K5"].value == "RFQ-000001-0001"
    assert sheet["A5"].value == "担当者：田中"
    assert sheet["A7"].value == "TEL:03-1234-5678"
    assert sheet["I12"].value == "船名【MITSUI OCEAN FUJI】"
    assert sheet["H15"].value is None


def test_tampered_source_workbook_is_rejected() -> None:
    template = _template()
    workbook = load_workbook(TEMPLATE_PATH)
    workbook.active["C21"] = "wrong heading"

    with pytest.raises(ValueError, match="TEMPLATE_CONTRACT_MISMATCH"):
        assert_source_workbook_matches(workbook.active, template)


def test_invalid_editable_contract_is_rejected() -> None:
    definition = standard_template_definition()
    definition["template_styles"]["zones"]["summary"]["start"] = 30
    definition["template_styles"]["product_row_formulas"]["L"] = "=H22*J22"

    errors = configuration_errors(
        field_positions=definition["field_positions"],
        product_table_config=definition["product_table_config"],
        template_styles=definition["template_styles"],
    )

    assert any("区域必须满足" in error for error in errors)
    assert any("包含 {row}" in error for error in errors)
