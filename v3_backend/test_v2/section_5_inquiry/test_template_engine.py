"""Section 5 — Inquiry: Excel rendering engine + template selector.

测试目标：
    供应商询价单的生成靠两件事——选模板 + 渲染模板。这两层都是纯函数（无 DB
    副作用、无 LLM、无 IO），所以我们用 in-memory ORM rows + openpyxl 直接读
    生成的字节，断言每一个 cell。

为什么重要：
    模板选错或者 cell 位置不对，供应商收到的询价单就是错的——历史 bug 已经
    在 v2 出现过（H8 写当前日期、I22 unit 写死成 CT、F 列 pack_size 用错值）。
    这里是给生产环境守门，确保 v3 重写后行为和 v2 等价。

设计方法：
    - render_inquiry_excel: 准备 SupplierTemplate ORM 对象 (in-memory)，喂订单
      字典 + 产品列表，调用渲染函数，从返回的 bytes 用 openpyxl 读出 cell。
    - resolve_template / select_template: 构造一组 SupplierTemplate ORM 对象，
      调用 resolver，断言返回的 (template, method, candidates) 三元组。
"""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from domains.inquiry.models import SupplierTemplate
from domains.inquiry.template_engine import render_inquiry_excel
from domains.inquiry.template_selector import (
    resolve_template,
    select_template,
    template_has_zone_config,
)

# ─── Helpers ──────────────────────────────────────────────────


def _zoned() -> dict:
    """Marker for a production template — selector只识别带 zones 的模板。"""
    return {"zones": {"meta": {}, "product_table": {}}}


def _make_template(
    *,
    id: int,
    name: str = "Default",
    supplier_id: int | None = None,
    supplier_ids: list[int] | None = None,
    country_id: int | None = None,
    field_positions: dict | None = None,
    product_table_config: dict | None = None,
    zoned: bool = True,
) -> SupplierTemplate:
    """Construct an in-memory SupplierTemplate. No DB needed for selector tests."""
    return SupplierTemplate(
        id=id,
        template_name=name,
        supplier_id=supplier_id,
        supplier_ids=supplier_ids,
        country_id=country_id,
        field_positions=field_positions,
        product_table_config=product_table_config,
        template_styles=_zoned() if zoned else None,
        has_product_table=True,
    )


def _open(xlsx_bytes: bytes):
    return load_workbook(io.BytesIO(xlsx_bytes))


# ─── render_inquiry_excel — generic fallback ──────────────────


def test_render_generic_when_no_template():
    """No template → generic header + meta block + product table."""
    metadata = {
        "po_number": "PO-2026-001",
        "order_date": "2026-05-01",
        "delivery_date": "2026-05-10",
        "currency": "USD",
        "ship_name": "MV Tester",
        "vendor_name": "Vendor X",
        "destination_port": "Yokohama",
    }
    products = [
        {
            "product_code": "X1",
            "product_name": "Apples",
            "quantity": 5,
            "unit_price": 2.0,
            "unit": "KG",
            "matched_product": {"code": "X1", "product_name_en": "Apples DB"},
        },
        {
            "product_code": "X2",
            "product_name": "Bananas",
            "quantity": 3,
            "unit_price": 1.5,
            "unit": "KG",
            "matched_product": {"code": "X2", "product_name_en": "Bananas DB"},
        },
    ]
    out = render_inquiry_excel(
        None, metadata=metadata, products=products, supplier_id=999
    )
    wb = _open(out)
    ws = wb.active
    # Title banner at A1 (merged)
    assert "Purchase Order" in str(ws["A1"].value)
    # Metadata cells filled in
    assert ws["B3"].value == "PO-2026-001"
    assert ws["B5"].value == "2026-05-10"
    assert ws["B7"].value == "USD"
    # Supplier id at E3
    assert ws["E3"].value == "999"
    # Header row at row 9
    assert ws.cell(row=9, column=1).value == "No."
    assert ws.cell(row=9, column=3).value == "Product Name"
    # Product rows start at row 10
    assert ws.cell(row=10, column=1).value == 1
    assert ws.cell(row=10, column=2).value == "X1"
    assert ws.cell(row=10, column=4).value == 5
    assert ws.cell(row=10, column=5).value == "KG"
    assert ws.cell(row=11, column=2).value == "X2"
    # Total row appears after the last product
    total_row = 9 + len(products) + 1
    assert ws.cell(row=total_row, column=5).value == "Total:"


def test_render_generic_uses_matched_product_when_extracted_missing():
    """If extracted product is missing fields, fall back to matched_product values."""
    metadata = {"po_number": "PO-1", "currency": "JPY"}
    products = [
        {
            # No product_code / product_name / unit_price on the extracted row
            "quantity": 2,
            "matched_product": {
                "code": "M-7",
                "product_name_en": "From DB",
                "price": 12.5,
                "unit": "L",
            },
        }
    ]
    out = render_inquiry_excel(
        None, metadata=metadata, products=products, supplier_id=1
    )
    ws = _open(out).active
    assert ws.cell(row=10, column=2).value == "M-7"
    assert ws.cell(row=10, column=3).value == "From DB"
    assert ws.cell(row=10, column=5).value == "L"
    assert ws.cell(row=10, column=6).value == 12.5


def test_render_generic_highlights_included_warning_row_yellow():
    products = [
        {
            "product_code": "W-1",
            "product_name": "Warning",
            "quantity": 1,
            "unit": "EA",
            "inquiry_warnings": [
                {"code": "SELLING_PRICE_DEVIATION", "message": "卖价偏差"}
            ],
            "matched_product": {
                "code": "W-1",
                "product_name_en": "Warning",
                "price": 2,
                "unit": "EA",
            },
        }
    ]

    ws = _open(
        render_inquiry_excel(None, metadata={}, products=products, supplier_id=1)
    ).active

    assert all(
        ws.cell(row=10, column=column).fill.fgColor.rgb.endswith("FFF2CC")
        for column in range(1, 8)
    )


# ─── render_inquiry_excel — template-driven ───────────────────


def test_render_with_template_fills_field_positions():
    """SupplierTemplate cell-map → exact cells filled at H8 / I22 / F / C positions."""
    template = _make_template(
        id=1,
        name="Concrete",
        supplier_id=100,
        field_positions={
            "po_number": "C8",
            "delivery_date": "H8",
            "ship_name": "B4",
        },
        product_table_config={
            "start_row": 22,
            "columns": {
                "C": "product_code",
                "F": "description",
                "I": "unit",
                "K": "quantity",
                "M": "unit_price",
            },
        },
    )
    metadata = {
        "po_number": "PO-2026-001",
        "delivery_date": "2026-05-10",
        "ship_name": "MV Tester",
    }
    products = [
        {
            "product_code": "P1",
            "quantity": 5,
            "unit_price": 745.0,  # customer PO price — v44: must NOT reach sheet
            "unit": "KG2.2",  # corrupted — matched_product.unit should win
            "matched_product": {
                "code": "DBC-001",
                "pack_size": "6-10ct/10kg",
                "unit": "KG",
                "price": 690.0,  # our catalog price — what suppliers must see
            },
        },
    ]
    out = render_inquiry_excel(
        template, metadata=metadata, products=products, supplier_id=100
    )
    ws = _open(out).active
    # Header positions
    assert ws["C8"].value == "PO-2026-001"
    assert ws["H8"].value == "2026-05-10"
    assert ws["B4"].value == "MV Tester"
    # Product row at start_row=22
    assert ws["C22"].value == "P1"
    # description → matched_product.pack_size (DB-authoritative)
    assert ws["F22"].value == "6-10ct/10kg"
    # unit → matched_product.unit wins over corrupted extracted unit
    assert ws["I22"].value == "KG"
    assert ws["K22"].value == 5
    # unit_price → v44 invariant: DB catalog price ONLY (690), customer's
    # PO quote (745) stays out. Repros prod 2026-05-20 order #104 incident.
    assert ws["M22"].value == 690.0
    assert ws["M22"].value != 745.0


def test_render_template_highlights_actual_warning_product_columns():
    template = _make_template(
        id=99,
        field_positions={"po_number": "A1"},
        product_table_config={
            "start_row": 22,
            "columns": {"C": "product_code", "K": "quantity"},
        },
        zoned=False,
    )
    products = [
        {
            "product_code": "W-2",
            "quantity": 3,
            "inquiry_warnings": [
                {"code": "SELLING_PRICE_DEVIATION", "message": "卖价偏差"}
            ],
            "matched_product": {"price": 2, "unit": "EA"},
        }
    ]

    ws = _open(
        render_inquiry_excel(template, metadata={}, products=products, supplier_id=1)
    ).active

    assert ws["C22"].fill.fgColor.rgb.endswith("FFF2CC")
    assert ws["K22"].fill.fgColor.rgb.endswith("FFF2CC")


def test_render_template_skips_formula_columns():
    """Columns listed in `formula_columns` are NOT overwritten by the engine."""
    template = _make_template(
        id=2,
        # `field_positions` must be truthy to take the template path (not generic).
        field_positions={"po_number": "A1"},
        product_table_config={
            "start_row": 10,
            "columns": {"D": "quantity", "E": "unit_price", "F": "total_price"},
            "formula_columns": ["F"],
        },
    )
    products = [{
        "quantity": 5,
        "unit_price": 2.0,
        "total_price": 999,
        # v44: unit_price comes from DB. Add matched_product.price so the
        # test still exercises a non-formula column write, not the empty
        # path (which has its own dedicated test below).
        "matched_product": {"price": 2.0},
    }]
    out = render_inquiry_excel(
        template, metadata={}, products=products, supplier_id=1
    )
    ws = _open(out).active
    assert ws["D10"].value == 5
    assert ws["E10"].value == 2.0
    # total_price is a formula column → engine must NOT have written 999.
    assert ws["F10"].value is None


def test_render_template_handles_unit_price_zero():
    """0.00 must be written (used to be dropped by `if not x` checks — v2 bug).

    v44: unit_price is sourced from DB only. We assert that a DB price of
    0 (free sample) reaches the sheet — the historical concern was that
    0 was indistinguishable from None and got dropped.
    """
    template = _make_template(
        id=3,
        # Non-empty field_positions forces the template-driven path.
        field_positions={"po_number": "Z1"},
        product_table_config={
            "start_row": 15,
            "columns": {"D": "unit_price"},
        },
    )
    products = [{"unit_price": 999, "matched_product": {"price": 0}}]
    out = render_inquiry_excel(
        template, metadata={}, products=products, supplier_id=1
    )
    ws = _open(out).active
    assert ws["D15"].value == 0


def test_render_template_unit_price_none_in_db_leaves_cell_empty():
    """v44 invariant: when DB has no price (matched_product.price is None),
    the unit_price cell is left empty — the Excel value is NOT used as a
    fallback. Buyer must fill in the catalog price before the inquiry
    sheet has a number to quote off."""
    template = _make_template(
        id=99,
        field_positions={"po_number": "Z1"},
        product_table_config={
            "start_row": 16,
            "columns": {"D": "unit_price"},
        },
    )
    # Excel has a price (745); DB doesn't. Cell must end up empty.
    products = [{"unit_price": 745.0, "matched_product": {"code": "X"}}]
    out = render_inquiry_excel(
        template, metadata={}, products=products, supplier_id=1
    )
    ws = _open(out).active
    assert ws["D16"].value is None or ws["D16"].value == ""


def test_render_template_unit_price_no_matched_product_leaves_cell_empty():
    """v44 boundary: when the product wasn't matched (matched_product is
    None / absent), the unit_price cell stays empty. There's no DB row to
    quote from, so we don't silently use the customer's PO price."""
    template = _make_template(
        id=98,
        field_positions={"po_number": "Z1"},
        product_table_config={
            "start_row": 17,
            "columns": {"D": "unit_price"},
        },
    )
    products = [{"unit_price": 745.0}]  # no matched_product
    out = render_inquiry_excel(
        template, metadata={}, products=products, supplier_id=1
    )
    ws = _open(out).active
    assert ws["D17"].value is None or ws["D17"].value == ""


def test_render_template_field_mapping_indirection():
    """`field_mapping` lets a template key (e.g. "po") map to a metadata key
    (e.g. "po_number") without changing the template config."""
    template = _make_template(
        id=4,
        field_positions={"po": "A1"},
        product_table_config={"start_row": 10, "columns": {}},
    )
    out = render_inquiry_excel(
        template,
        metadata={"po_number": "PO-XYZ"},
        products=[],
        supplier_id=1,
        field_mapping={"po": "po_number"},
    )
    ws = _open(out).active
    assert ws["A1"].value == "PO-XYZ"


def test_render_template_products_grouped_into_consecutive_rows():
    """Three products → rows start_row, start_row+1, start_row+2."""
    template = _make_template(
        id=5,
        # Non-empty field_positions forces the template-driven path.
        field_positions={"po_number": "Z1"},
        product_table_config={
            "start_row": 22,
            "columns": {"D": "line_number", "E": "product_code"},
        },
    )
    products = [
        {"product_code": "AAA"},
        {"product_code": "BBB"},
        {"product_code": "CCC"},
    ]
    out = render_inquiry_excel(
        template, metadata={}, products=products, supplier_id=1
    )
    ws = _open(out).active
    assert ws["D22"].value == 1
    assert ws["E22"].value == "AAA"
    assert ws["D23"].value == 2
    assert ws["E23"].value == "BBB"
    assert ws["D24"].value == 3
    assert ws["E24"].value == "CCC"


# ─── template_has_zone_config ─────────────────────────────────


def test_template_has_zone_config_true_when_zones_dict_present():
    tpl = _make_template(id=1, zoned=True)
    assert template_has_zone_config(tpl) is True


def test_template_has_zone_config_false_for_legacy_template():
    tpl = _make_template(id=1, zoned=False)
    assert template_has_zone_config(tpl) is False


def test_template_has_zone_config_handles_none():
    assert template_has_zone_config(None) is False


# ─── resolve_template ─────────────────────────────────────────


def test_resolve_template_exact_supplier_ids_match():
    """`supplier_ids` array containing the target → method="exact"."""
    a = _make_template(id=1, name="A", supplier_ids=[100, 200])
    b = _make_template(id=2, name="B", supplier_id=999)
    tpl, method, cands = resolve_template(100, [a, b])
    assert tpl is a
    assert method == "exact"
    assert cands == []


def test_resolve_template_legacy_single_supplier_id_match():
    """`supplier_id` scalar (old shape) is still respected."""
    a = _make_template(id=1, name="A", supplier_ids=[])
    b = _make_template(id=2, name="B", supplier_id=42)
    tpl, method, cands = resolve_template(42, [a, b])
    assert tpl is b
    # Production code returns "exact" for either path.
    assert method == "exact"
    assert cands == []


def test_resolve_template_returns_candidates_when_no_match():
    """No exact binding → caller must pick from candidates."""
    a = _make_template(id=1, name="A", supplier_id=100, country_id=11)
    b = _make_template(id=2, name="B", supplier_id=200, country_id=22)
    tpl, method, cands = resolve_template(999, [a, b])
    assert tpl is None
    assert method == "candidates"
    assert len(cands) == 2
    ids = {c["id"] for c in cands}
    assert ids == {1, 2}
    # Candidates carry country_id so the UI can group by country.
    countries = {c["country_id"] for c in cands}
    assert countries == {11, 22}


def test_resolve_template_returns_unavailable_when_no_templates():
    tpl, method, cands = resolve_template(1, [])
    assert tpl is None
    assert method == "unavailable"
    assert cands == []


def test_resolve_template_ignores_legacy_templates_without_zones():
    """Templates without `template_styles.zones` are filtered out."""
    legacy = _make_template(id=1, supplier_ids=[100], zoned=False)
    tpl, method, cands = resolve_template(100, [legacy])
    # Legacy template is excluded from the production pool → unavailable.
    assert tpl is None
    assert method == "unavailable"


# ─── select_template (auto-pick) ──────────────────────────────


def test_select_template_auto_picks_first_candidate():
    """No exact match → select_template returns "candidate_auto" + first one."""
    a = _make_template(id=1, name="A", supplier_id=100)
    b = _make_template(id=2, name="B", supplier_id=200)
    tpl, method, cands = select_template(999, [a, b])
    assert tpl is not None
    assert method == "candidate_auto"
    assert tpl.id in {1, 2}
    # Candidates are still surfaced so the caller can show alternatives.
    assert len(cands) == 2


def test_select_template_user_override_returns_user_selected():
    a = _make_template(id=1, name="A", supplier_id=100)
    b = _make_template(id=2, name="B", supplier_id=200)
    tpl, method, _cands = select_template(100, [a, b], template_id_override=2)
    assert tpl is b
    assert method == "user_selected"


def test_select_template_unknown_override_id_raises():
    a = _make_template(id=1, supplier_id=100)
    with pytest.raises(ValueError, match="不存在"):
        select_template(100, [a], template_id_override=42)


def test_select_template_legacy_override_raises():
    """User can't manually force a legacy (non-zoned) template."""
    legacy = _make_template(id=7, supplier_id=100, zoned=False)
    with pytest.raises(ValueError, match="已下架"):
        select_template(100, [legacy], template_id_override=7)


def test_select_template_returns_unavailable_when_no_production_pool():
    """Empty pool + no override → ("None", "unavailable", [])."""
    legacy = _make_template(id=1, supplier_id=100, zoned=False)
    tpl, method, cands = select_template(100, [legacy])
    assert tpl is None
    assert method == "unavailable"
    assert cands == []


# ─── Polluted-template + type-coercion (2026-05 regression) ──────────
#
# Production hit a #VALUE!-everywhere bug: every supplier's inquiry sheet
# shipped 83 product rows with Sub Total / Tax / GRAND TOTAL all #VALUE!.
# Root cause: the uploaded SupplierTemplate carried 83 rows of *example*
# product data left over from a real order. We overwrote only the first N
# rows and left the rest as-is — one of which had a string ("確認中" = TBD)
# in the quantity cell, so the per-row `=H*K` formula exploded and SUM
# range propagated it. These tests pin the two fixes:
#   1. strip pre-filled rows from the template before writing real data
#   2. coerce string-quantity to None instead of writing it raw


from domains.inquiry.template_engine import (  # noqa: E402 — grouped with regression tests
    _clear_data_cells,
    _detect_filled_data_rows,
    _to_number_or_none,
)

# --- _to_number_or_none (pure) -----------------------------------------


def test_to_number_or_none_passes_numbers_through():
    assert _to_number_or_none(5) == 5
    assert _to_number_or_none(5.5) == 5.5
    assert _to_number_or_none(0) == 0  # falsy-but-valid — never drop


def test_to_number_or_none_parses_numeric_strings():
    assert _to_number_or_none("100") == 100.0
    assert _to_number_or_none("100.5") == 100.5
    # Stray whitespace from PDF extraction must not block coercion.
    assert _to_number_or_none(" 42 ") == 42.0


def test_to_number_or_none_drops_non_numeric_strings():
    """`確認中` is the exact placeholder string that triggered the prod bug."""
    assert _to_number_or_none("確認中") is None
    assert _to_number_or_none("TBD") is None
    assert _to_number_or_none("") is None
    assert _to_number_or_none("   ") is None


def test_to_number_or_none_drops_unsupported_types():
    assert _to_number_or_none(None) is None
    assert _to_number_or_none([1, 2]) is None
    assert _to_number_or_none({"a": 1}) is None


# --- _detect_filled_data_rows / _clear_data_cells (operate on ws) ------


def _polluted_template_ws():
    """Build a Workbook whose product range (rows 19..21) is pre-filled —
    the exact pattern the production template shipped with."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    # 3 pre-filled sample product rows
    for r, (code, qty, price) in enumerate(
        [("X1", 10, 50), ("X2", 20, 60), ("X3", 30, 70)], start=19
    ):
        ws.cell(row=r, column=3, value=code)  # C: product_code
        ws.cell(row=r, column=8, value=qty)  # H: quantity
        ws.cell(row=r, column=11, value=price)  # K: unit_price
        ws.cell(row=r, column=13, value=f"=H{r}*K{r}")  # M: formula (kept)
    return wb, ws


def test_detect_filled_data_rows_returns_last_filled():
    columns = {"C": "product_code", "H": "quantity", "K": "unit_price", "M": "amount"}
    formula = {"M"}
    _, ws = _polluted_template_ws()
    last = _detect_filled_data_rows(ws, 19, columns, formula)
    assert last == 21


def test_detect_filled_data_rows_returns_start_minus_one_when_empty():
    """Empty template → caller skips the clear pass."""
    from openpyxl import Workbook

    ws = Workbook().active
    last = _detect_filled_data_rows(ws, 19, {"C": "x", "H": "y"}, set())
    assert last == 18  # start_row - 1


def test_detect_filled_data_rows_ignores_formula_columns():
    """A row whose ONLY filled cell is a formula column shouldn't count as
    sample data — formulas are template machinery."""
    from openpyxl import Workbook

    ws = Workbook().active
    ws["M19"] = "=H19*K19"  # only formula present
    last = _detect_filled_data_rows(ws, 19, {"H": "qty", "M": "amount"}, {"M"})
    assert last == 18


def test_clear_data_cells_blanks_non_formula_cells():
    wb, ws = _polluted_template_ws()
    columns = {"C": "product_code", "H": "quantity", "K": "unit_price", "M": "amount"}
    _clear_data_cells(ws, 19, 21, columns, {"M"})

    for r in range(19, 22):
        assert ws.cell(row=r, column=3).value is None  # C cleared
        assert ws.cell(row=r, column=8).value is None  # H cleared
        assert ws.cell(row=r, column=11).value is None  # K cleared
        # M (formula) is preserved — Excel will recalc on open.
        assert ws.cell(row=r, column=13).value == f"=H{r}*K{r}"


# --- _fill_with_template integration (regression for the #VALUE! bug) --


def _polluted_template_bytes():
    """Materialise the polluted template as bytes the engine can load_workbook on."""
    wb, _ = _polluted_template_ws()
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_render_strips_pre_filled_sample_rows():
    """1 real product against a 3-row polluted template → exactly 1 data row;
    rows 20-21 must be wiped clean (no leftover sample data)."""
    template = _make_template(
        id=1,
        field_positions={"po_number": "A1"},
        product_table_config={
            "start_row": 19,
            "columns": {
                "C": "product_code",
                "H": "quantity",
                "K": "unit_price",
            },
            "formula_columns": ["M"],
        },
    )
    out = render_inquiry_excel(
        template,
        metadata={"po_number": "PO-CLEAN"},
        # v44: unit_price source = matched_product.price. Customer PO price
        # in `unit_price` field is no longer used by the engine.
        products=[{
            "product_code": "REAL",
            "quantity": 7,
            "unit_price": 999,  # would-be Excel value (ignored)
            "matched_product": {"price": 99},
        }],
        supplier_id=1,
        template_file_bytes=_polluted_template_bytes(),
    )
    ws = _open(out).active
    # Real product landed at start_row.
    assert ws["C19"].value == "REAL"
    assert ws["H19"].value == 7
    assert ws["K19"].value == 99
    # Sample rows 20 & 21 must be clean. Pre-fix they would still contain
    # X2/20/60 and X3/30/70.
    for r in (20, 21):
        assert ws.cell(row=r, column=3).value is None
        assert ws.cell(row=r, column=8).value is None
        assert ws.cell(row=r, column=11).value is None
    # Formula column survives — Excel recalcs on open.
    assert ws.cell(row=19, column=13).value == "=H19*K19"


def test_render_coerces_string_quantity_to_blank_not_value_bug():
    """quantity='確認中' must NOT land in the cell — leaving it would make
    Excel's `=H*K` formula return #VALUE!, which cascades into Sub Total,
    Tax, GRAND TOTAL, and the in-header TOTAL via `=M_grand`. Pre-fix this
    string was written raw; post-fix the cell stays blank (Excel reads
    blank as 0)."""
    template = _make_template(
        id=2,
        field_positions={"po_number": "A1"},
        product_table_config={
            "start_row": 19,
            "columns": {"H": "quantity", "K": "unit_price"},
            "formula_columns": [],
        },
    )
    out = render_inquiry_excel(
        template,
        metadata={"po_number": "PO"},
        # v44: unit_price source = matched_product.price (not the raw
        # `unit_price` field, which holds the customer's PO price).
        products=[{
            "quantity": "確認中",
            "unit_price": 9999,
            "matched_product": {"price": 5000},
        }],
        supplier_id=1,
        template_file_bytes=_polluted_template_bytes(),
    )
    ws = _open(out).active
    # The string was coerced out — cell is blank, formula sees 0.
    assert ws["H19"].value is None
    # But unit_price is a normal number → written.
    assert ws["K19"].value == 5000


def test_render_keeps_zero_quantity_as_real_zero():
    """Edge: quantity=0 is a *valid* number, not a placeholder. It must
    survive coercion (otherwise we'd treat a real 0-qty line as missing data)."""
    template = _make_template(
        id=3,
        field_positions={"po_number": "A1"},
        product_table_config={
            "start_row": 19,
            "columns": {"H": "quantity"},
            "formula_columns": [],
        },
    )
    out = render_inquiry_excel(
        template,
        metadata={"po_number": "PO"},
        products=[{"quantity": 0}],
        supplier_id=1,
        template_file_bytes=_polluted_template_bytes(),
    )
    ws = _open(out).active
    assert ws["H19"].value == 0  # not None, not skipped
