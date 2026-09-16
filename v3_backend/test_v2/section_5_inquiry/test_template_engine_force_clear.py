"""Section 5 — inquiry: template engine MUST force-clear field cells.

测试目标：
    Pin the render-time behaviour introduced 2026-05-22.
    Prod bug: supplier-uploaded template files carry stale sample data in
    field_positions cells (e.g. I8 = a real Tokyo address baked in when the
    supplier rendered a previous order). The old renderer used `if value ==
    "": continue` and so stale samples leaked through to every customer.

    Contract:
      • Every cell declared in template.field_positions is wiped to None first.
      • Cells with non-empty metadata values then get the new value written.
      • Cells whose metadata value is empty/None stay blank — never reverted
        to whatever the template had baked in.

为什么必须有：
    Without this test, anyone editing `_fill_with_template` can re-introduce
    the leak in a single line change. The bug is invisible to anyone who
    only tests with metadata that happens to provide every field.
"""

from __future__ import annotations

import io

from openpyxl import Workbook, load_workbook

from domains.inquiry.models import SupplierTemplate
from domains.inquiry.template_engine import render_inquiry_excel


def _zoned() -> dict:
    return {"zones": {"meta": {}, "product_table": {}}}


def _make_template_with_stale_xlsx() -> tuple[SupplierTemplate, bytes]:
    """A template file with sample data pre-filled in field_positions cells —
    simulates a supplier who saved a real-order Excel as their template."""
    wb = Workbook()
    ws = wb.active
    ws["A3"] = "STALE_SUPPLIER_NAME"
    ws["B3"] = "STALE_PO"
    ws["I8"] = "〒104-0053 東京都中央区晴海5-7-1\n船名【シルバーシー・ノバ】"
    ws["I9"] = "STALE_SHIP_JP"
    buf = io.BytesIO()
    wb.save(buf)
    template = SupplierTemplate(
        id=99,
        template_name="日本订单标准",
        field_positions={
            "supplier_name": "A3",
            "po_number": "B3",
            "delivery_address": "I8",
            "ship_name_jp": "I9",
        },
        product_table_config={"columns": {}, "start_row": 12},
        template_styles=_zoned(),
        has_product_table=True,
    )
    return template, buf.getvalue()


def _open(xlsx_bytes: bytes):
    return load_workbook(io.BytesIO(xlsx_bytes))


def test_force_clear_wipes_stale_cells_when_metadata_missing_field() -> None:
    """Metadata has no delivery_address → I8 must be blank, NOT the stale sample."""
    template, file_bytes = _make_template_with_stale_xlsx()
    metadata = {
        "supplier_name": "REAL Supplier",
        "po_number": "PO-REAL-001",
        # delivery_address intentionally absent
        # ship_name_jp intentionally absent
    }
    out = render_inquiry_excel(
        template,
        metadata=metadata,
        products=[],
        supplier_id=1,
        template_file_bytes=file_bytes,
        field_mapping=None,
    )
    ws = _open(out).active
    assert ws["A3"].value == "REAL Supplier"
    assert ws["B3"].value == "PO-REAL-001"
    # The whole point of the fix — these MUST be empty, not the stale samples.
    assert ws["I8"].value is None, (
        f"I8 should be cleared, got {ws['I8'].value!r}"
    )
    assert ws["I9"].value is None, (
        f"I9 should be cleared, got {ws['I9'].value!r}"
    )


def test_force_clear_then_writes_when_metadata_provides_value() -> None:
    """Metadata has delivery_address → I8 must show the NEW value, not the stale one."""
    template, file_bytes = _make_template_with_stale_xlsx()
    metadata = {
        "supplier_name": "REAL Supplier",
        "po_number": "PO-REAL-002",
        "delivery_address": "3-11-8 Chikko, Minato-ku, Osaka City",
        "ship_name_jp": "CELEBRITY MILLENNIUM",
    }
    out = render_inquiry_excel(
        template,
        metadata=metadata,
        products=[],
        supplier_id=1,
        template_file_bytes=file_bytes,
        field_mapping=None,
    )
    ws = _open(out).active
    assert ws["I8"].value == "3-11-8 Chikko, Minato-ku, Osaka City"
    assert ws["I9"].value == "CELEBRITY MILLENNIUM"
