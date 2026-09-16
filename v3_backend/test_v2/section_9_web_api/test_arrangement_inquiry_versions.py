"""Arrangement inquiries are immutable, source-aware historical versions."""

from datetime import datetime, timedelta
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from domains.inquiry import repository as inquiry_repo
from domains.inquiry.models import SupplierTemplate
from domains.inquiry.orchestrator import queue_inquiry_for_group, run_inquiry_for_group
from domains.masterdata.models import Country, Port, Product, Supplier
from domains.masterdata.price_periods import create_period
from domains.masterdata.schemas import ProductPricePeriodCreate
from domains.orders.models import Order, OrderGroup
from infrastructure.storage import get_storage
from test_v2.fixtures.helpers import login, seed_user


def _scope(db):
    user = seed_user(db, email="arrangement-version@test")
    country = Country(name="Japan", code="JPN")
    db.add(country)
    db.flush()
    port = Port(name="横浜", code="YOK", country_id=country.id)
    supplier = Supplier(name="供应商 A")
    db.add_all([port, supplier])
    db.flush()
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "PO Number"
    sheet["A2"] = "Ship Name"
    sheet["A5"] = "Code"
    sheet["B5"] = "Product"
    sheet["C5"] = "Quantity"
    buffer = BytesIO()
    workbook.save(buffer)
    template_url = get_storage().upload(
        "templates",
        f"arrangement-{user.id}.xlsx",
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    db.add(SupplierTemplate(
        supplier_id=supplier.id,
        template_name="安排标准模板",
        template_file_url=template_url,
        field_positions={
            "po_number": {"position": "B1"},
            "ship_name": {"position": "B2"},
        },
        product_table_config={
            "start_row": 6,
            "columns": {"A": "product_code", "B": "product_name_en", "C": "quantity"},
        },
        template_styles={"zones": {"product_data": {"start": 6, "end": 50}}},
    ))
    group = OrderGroup(
        user_id=user.id,
        name="2026-10-01 · 横浜",
        loading_date="2026-10-01",
    )
    db.add(group)
    db.flush()
    product = Product(
        code="APPLE",
        product_name_en="Apple",
        country_id=country.id,
        port_id=port.id,
        supplier_id=supplier.id,
        unit="CA",
        price=30,
        status=True,
    )
    db.add(product)
    db.commit()
    return user, port, supplier, group


def _add_order(db, *, user_id, group_id, port_id, po, code="APPLE"):
    order = Order(
        user_id=user_id,
        group_id=group_id,
        port_id=port_id,
        filename=f"{po}.pdf",
        po_number=po,
        loading_date="2026-10-01",
        currency="JPY",
        products=[
            {
                "line_id": "line-00001",
                "source_line": "1",
                "page": 1,
                "product_code": code,
                "product_name": "Apple",
                "quantity": "2",
                "unit": "CA",
            }
        ],
        status="ready",
    )
    db.add(order)
    db.commit()
    db.refresh(order)
    return order


def test_each_run_keeps_history_and_defaults_to_latest(db):
    user, port, _supplier, group = _scope(db)
    anchor = _add_order(
        db,
        user_id=user.id,
        group_id=group.id,
        port_id=port.id,
        po="PO-A",
    )

    first = run_inquiry_for_group(group.id, max_workers=1)
    _add_order(
        db,
        user_id=user.id,
        group_id=group.id,
        port_id=port.id,
        po="PO-B",
    )
    second = run_inquiry_for_group(group.id, max_workers=1)

    assert first.group_id == second.group_id == group.id
    assert (first.version, second.version) == (1, 2)
    assert len(first.member_snapshot or []) == 1
    assert len(second.member_snapshot or []) == 2
    versions = inquiry_repo.list_inquiries_by_group(db, group.id)
    assert [row.version for row in versions] == [2, 1]
    assert inquiry_repo.get_inquiry_by_order(db, anchor.id).id == second.id


def test_queued_version_is_persistent_and_duplicate_clicks_coalesce(db):
    user, port, _supplier, group = _scope(db)
    _add_order(
        db,
        user_id=user.id,
        group_id=group.id,
        port_id=port.id,
        po="PO-QUEUED",
    )

    first = queue_inquiry_for_group(db, group.id)
    duplicate = queue_inquiry_for_group(db, group.id)

    assert first.id == duplicate.id
    assert first.status == "pending"
    completed = run_inquiry_for_group(group.id, inquiry_id=first.id, max_workers=1)
    assert completed.status == "completed"
    assert completed.run_attempts == 1
    next_version = queue_inquiry_for_group(db, group.id)
    assert next_version.id != first.id
    assert next_version.version == 2


def test_unexpected_failure_schedules_same_version_for_retry(db, monkeypatch):
    user, port, _supplier, group = _scope(db)
    _add_order(
        db,
        user_id=user.id,
        group_id=group.id,
        port_id=port.id,
        po="PO-RETRY",
    )
    queued = queue_inquiry_for_group(db, group.id)

    def crash(*_args, **_kwargs):
        raise RuntimeError("temporary storage outage")

    monkeypatch.setattr("domains.orders.service.match_order_group", crash)
    with pytest.raises(RuntimeError, match="temporary storage outage"):
        run_inquiry_for_group(group.id, inquiry_id=queued.id, max_workers=1)

    db.expire_all()
    failed = inquiry_repo.get_inquiry(db, queued.id)
    assert failed.status == "pending"
    assert failed.run_attempts == 1
    assert failed.next_retry_at is not None
    assert failed.error_message == "temporary storage outage"


def test_recovery_requeues_stale_running_version(db, monkeypatch):
    from apps.jobs import inquiry_jobs

    user, port, _supplier, group = _scope(db)
    _add_order(
        db,
        user_id=user.id,
        group_id=group.id,
        port_id=port.id,
        po="PO-RECOVER",
    )
    queued = queue_inquiry_for_group(db, group.id)
    row = inquiry_repo.get_inquiry(db, queued.id)
    row.status = "in_progress"
    row.heartbeat_at = datetime.utcnow() - timedelta(minutes=20)
    db.commit()
    submitted = []
    monkeypatch.setattr(
        inquiry_jobs,
        "submit_group_inquiry",
        lambda **kwargs: submitted.append(kwargs) or f"job-{kwargs['inquiry_id']}",
    )

    job_ids = inquiry_jobs.resume_pending_group_inquiries()

    assert job_ids == [f"job-{queued.id}"]
    assert submitted == [{"group_id": group.id, "inquiry_id": queued.id}]
    db.expire_all()
    recovered = inquiry_repo.get_inquiry(db, queued.id)
    assert recovered.status == "pending"
    assert recovered.error_message == "上次执行中断，系统已自动恢复"


def test_partial_version_generates_matched_file_and_keeps_unmatched(db):
    user, port, _supplier, group = _scope(db)
    _add_order(
        db,
        user_id=user.id,
        group_id=group.id,
        port_id=port.id,
        po="PO-MATCHED",
    )
    _add_order(
        db,
        user_id=user.id,
        group_id=group.id,
        port_id=port.id,
        po="PO-UNMATCHED",
        code="MISSING",
    )

    state = run_inquiry_for_group(group.id, max_workers=1)

    assert state.status == "partial"
    assert state.unassigned_count == 1
    assert state.unmatched_items[0]["source_po_number"] == "PO-UNMATCHED"
    assert state.unmatched_items[0]["match_reason"] == "未找到匹配商品"
    assert len(state.suppliers) == 1
    assert state.suppliers[0].status == "completed"


def test_generated_workbook_identifies_each_source_po(db):
    user, port, _supplier, group = _scope(db)
    for po in ("PO-A", "PO-B"):
        _add_order(
            db,
            user_id=user.id,
            group_id=group.id,
            port_id=port.id,
            po=po,
        )

    state = run_inquiry_for_group(group.id, max_workers=1)
    content = get_storage().download(state.suppliers[0].excel_file_url)
    workbook = load_workbook(BytesIO(content), data_only=False)
    source = workbook["PO Source"]
    values = {
        str(cell.value)
        for row in source.iter_rows()
        for cell in row
        if cell.value is not None
    }

    assert {"PO-A", "PO-B"}.issubset(values)


def test_arrangement_generation_never_silently_uses_generic_workbook(db):
    user, port, _supplier, group = _scope(db)
    db.query(SupplierTemplate).delete()
    db.commit()
    _add_order(
        db,
        user_id=user.id,
        group_id=group.id,
        port_id=port.id,
        po="PO-NO-TEMPLATE",
    )

    state = run_inquiry_for_group(group.id, max_workers=1)

    assert state.status == "error"
    assert state.suppliers[0].status == "error"
    assert "未配置可用的询价模板" in state.suppliers[0].error_message


def test_history_endpoint_returns_newest_first(client, db):
    user, port, _supplier, group = _scope(db)
    headers = login(client, user.email)
    _add_order(
        db,
        user_id=user.id,
        group_id=group.id,
        port_id=port.id,
        po="PO-A",
    )
    run_inquiry_for_group(group.id, max_workers=1)
    run_inquiry_for_group(group.id, max_workers=1)

    response = client.get(f"/api/order-groups/{group.id}/inquiries", headers=headers)

    assert response.status_code == 200, response.text
    assert [row["version"] for row in response.json()] == [2, 1]


def test_configured_purchase_price_gap_excludes_only_that_line(db):
    from datetime import date

    user, port, _supplier, group = _scope(db)
    product = db.query(Product).filter_by(code="APPLE").one()
    create_period(
        db,
        product.id,
        ProductPricePeriodCreate(
            price_type="purchase",
            amount=20,
            effective_from=date(2026, 1, 1),
            effective_to=date(2026, 6, 30),
        ),
        actor_id=None,
    )
    _add_order(
        db,
        user_id=user.id,
        group_id=group.id,
        port_id=port.id,
        po="PO-GAP",
    )

    state = run_inquiry_for_group(group.id, max_workers=1)

    assert state.status == "unmatched"
    assert state.unmatched_items[0]["inquiry_eligibility"] == "excluded"
    assert "未命中采购价期间" in state.unmatched_items[0]["match_reason"]
