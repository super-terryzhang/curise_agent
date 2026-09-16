"""Excel parsing endpoints — `/api/excel/*`.

Used by the settings page when uploading templates. Both endpoints accept an
.xlsx file and return either header detection (for order-template column
mapping) or all non-empty cell positions (for supplier-template field mapping).
"""

from __future__ import annotations

import hashlib
from io import BytesIO
from typing import Any

from fastapi import APIRouter, File, HTTPException, UploadFile, status
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from apps.http._deps import Admin

router = APIRouter(prefix="/excel", tags=["excel"])


@router.post("/parse")
async def parse_excel(
    _admin: Admin,
    file: UploadFile = File(...),
) -> dict[str, Any]:
    """Detect headers + sample rows + per-sheet fingerprint.

    Mirrors v2 `routes/excel.py::parse` shape so the settings page UI works.
    """
    content = await _read_xlsx(file)
    wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        sheets: list[dict[str, Any]] = []
        for ws in wb.worksheets:
            header_row_idx = _detect_header_row(ws)
            headers: list[dict[str, Any]] = []
            for cell in ws[header_row_idx] if ws.max_row else []:
                if cell.column is None or cell.value is None:
                    continue
                col_letter = get_column_letter(cell.column)
                headers.append({"column": col_letter, "label": str(cell.value).strip()})

            data_start = header_row_idx + 1
            sample_rows: list[list[str]] = []
            max_row = ws.max_row or data_start
            for row in ws.iter_rows(min_row=data_start, max_row=min(data_start + 4, max_row)):
                sample_rows.append([str(c.value) if c.value is not None else "" for c in row])

            fingerprint = _fingerprint([h["label"] for h in headers])
            sheets.append(
                {
                    "name": ws.title,
                    "headers": headers,
                    "header_row": header_row_idx,
                    "data_start_row": data_start,
                    "sample_rows": sample_rows,
                    "total_rows": ws.max_row or 0,
                    "fingerprint": fingerprint,
                }
            )
        return {"sheets": sheets}
    finally:
        wb.close()


@router.post("/parse-cells")
async def parse_excel_cells(
    _admin: Admin,
    file: UploadFile = File(...),
) -> dict[str, Any]:
    """Return all non-empty cells (capped at 50 rows per sheet) for template mapping."""
    content = await _read_xlsx(file)
    wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        sheets: list[dict[str, Any]] = []
        for ws in wb.worksheets:
            cells: list[dict[str, Any]] = []
            for row in ws.iter_rows(max_row=min(ws.max_row or 50, 50)):
                for cell in row:
                    if cell.value is None:
                        continue
                    col_letter = get_column_letter(cell.column)
                    cells.append(
                        {
                            "position": f"{col_letter}{cell.row}",
                            "value": str(cell.value).strip(),
                            "row": cell.row,
                            "col": col_letter,
                        }
                    )
            sheets.append({"name": ws.title, "cells": cells})
        return {"sheets": sheets}
    finally:
        wb.close()


# ─── helpers ──────────────────────────────────────────────────


async def _read_xlsx(file: UploadFile) -> bytes:
    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xls")):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "请上传 .xlsx 文件")
    content = await file.read()
    if not content:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "文件为空")
    return content


def _detect_header_row(ws: Any) -> int:
    """First row with ≥ 3 non-empty text cells, defaulting to 1."""
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 20, 20)):
        text_count = sum(
            1
            for cell in row
            if cell.value is not None and isinstance(cell.value, str) and cell.value.strip()
        )
        if text_count >= 3:
            return int(row[0].row)
    return 1


def _fingerprint(header_labels: list[str]) -> str:
    normalized = "|".join(sorted(h.lower().strip() for h in header_labels if h))
    return hashlib.sha256(normalized.encode()).hexdigest()[:16]
